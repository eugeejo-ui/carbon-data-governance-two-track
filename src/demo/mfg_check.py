"""09 제조 시연 검산 (클로드 코드 1회차)

docs/progress/09_demo_mfg.md 4-8-10의 검산 17개를 실행하고 결과를 출력함
기대값은 4-8-9 표를 코드 안 상수로 옮겨 적음 (생성 코드를 불러오지 않음 — 독립 검산)
"""
import csv
import sys
from decimal import ROUND_HALF_UP, Decimal, getcontext
from pathlib import Path

from openpyxl import load_workbook

# 독립 검산 규칙 (사용자 확정 2026-10-12): src/demo 파일을 불러오지 않음, 숫자는 문자 그대로 Decimal,
# 반올림은 ROUND_HALF_UP(0에서 먼 쪽), 정밀도 50자리
getcontext().prec = 50

ROOT = Path(__file__).resolve().parents[2]
DB = ROOT / "data" / "demo" / "mfg" / "db"
EXT = ROOT / "data" / "demo" / "mfg" / "ext"
MONTHS = [f"2025-{m:02d}" for m in range(1, 13)]
MAX_FAIL_LINES = 20

# ── 기대값 (09 4-8-9·4-8-10) ──────────────────────────────────────────────

ANCHOR_EMISSIONS = {"KG": "402537", "TW": "151271", "HBG": "64871", "SA": "23731",
                    "HI": "6178096"}
ANCHOR_ENERGY = {"KG": "8685", "TW": "2753", "HBG": "1420", "SA": "475", "HI": "41976"}
COMPANY_PLANTS = {"KG": ["KG_DJ", "KG_IC"], "TW": ["TW_HJ", "TW_NS"],
                  "HBG": ["HBG_CW", "HBG_GJ", "HBG_SA"], "SA": ["SA_PS"],
                  "HI": ["HI_CEM", "HI_SEC"]}
FUEL_FIRST_COMPANIES = {"KG", "TW", "SA"}
# 4-8-4 상수 (gov_factors 값을 다시 적음)
MJ = Decimal("9.6")
NET = Decimal("38.5") / Decimal("42.7")
GAS = Decimal("56.1")
TOL_07 = Decimal("0.0065")
TOL_08 = Decimal("0.06")
EXPECT_09 = {"KG": "0.4312", "TW": "0.4739", "HBG": "0.4297", "SA": "0.4764", "HI": "0.4230"}
EXPECT_10 = {"원산지 증빙 없을 때": "20967954.892", "한국산 증빙 시": "5746411.636"}
# 4-8-6 생산 표 annual_public (사업보고서 단위를 t·㎥로 바꾼 값)
ANNUAL_PUBLIC = {"1": "1606000", "2": "601000", "3": "180710", "4": "68165", "5": "239432",
                 "6": "142170", "7": None, "8": "15457", "9": "118955", "10": "9260000",
                 "11": "2022000", "12": "3222000"}
# #05 추가 (사용자 확정 2026-10-12): 표 2 머리 설명 · 4-8-6 배출 파일 빈칸 규칙
DIRECT_EQ_FUEL = ["KG_DJ", "KG_IC", "SA_PS", "TW_NS"]
DIRECT_SPLIT_BLANK = ["HBG_CW", "HBG_GJ", "HBG_SA", "HI_CEM", "HI_SEC"]
# #13 (사용자 확정 2026-10-12): 양 칸 9개만 더함, 빈칸은 0, 허용 오차 0
BANK_SUM_ROWS = {"KG_ALL": ["KG_DJ", "KG_IC"], "HI_ALL": ["HI_CEM", "HI_SEC"],
                 "SA_ALL": ["SA_PS"]}
BANK_ADDITIVE = ["direct_t", "indirect_t", "total_t", "fuel_tj", "elec_mwh",
                 "precursor_unknown_t", "precursor_kr_t", "eu_embedded_unknown_t",
                 "eu_embedded_kr_t"]
BANK_ANCHOR = {"KG_ALL": "402537", "HI_ALL": "6178096", "SA_ALL": "23731"}

# 표 1. 공장 연간 값 (12달 합)
TABLE_1 = """
| plant_id | fuel_tj_gross | elec_tj | energy_tj | Σ월 elec_mwh | direct_fuel_t | process_t | direct_t | indirect_t | total_t | Σ월 direct_fuel_alt_net_t |
|---|---|---|---|---|---|---|---|---|---|---|
| KG_DJ | 1,600.515 | 4,719.426 | 6,319.941 | 491,606.876 | 80,957.183 | 0.000 | 80,957.183 | 211,962.809 | 292,919.992 | 89,788.892 |
| KG_IC | 598.947 | 1,766.112 | 2,365.059 | 183,970.001 | 30,295.932 | 0.000 | 30,295.932 | 79,321.076 | 109,617.008 | 33,600.930 |
| TW_HJ | 225.466 | 945.466 | 1,170.932 | 98,486.041 | 11,404.514 | 14,456.800 | 25,861.314 | 46,669.370 | 72,530.684 | 12,648.643 |
| TW_NS | 530.407 | 1,051.661 | 1,582.068 | 109,548.023 | 26,829.029 | 0.000 | 26,829.029 | 51,911.287 | 78,740.316 | 29,755.834 |
| HBG_CW | 248.312 | 1,155.973 | 1,404.285 | 120,413.854 | 빈칸 | 빈칸 | 12,406.171 | 51,746.908 | 64,153.079 | 빈칸 |
| HBG_GJ | 2.321 | 10.803 | 13.124 | 1,125.312 | 빈칸 | 빈칸 | 115.944 | 483.610 | 599.554 | 빈칸 |
| HBG_SA | 0.458 | 2.133 | 2.591 | 222.188 | 빈칸 | 빈칸 | 22.890 | 95.477 | 118.367 | 빈칸 |
| SA_PS | 164.842 | 310.158 | 475.000 | 32,308.127 | 8,338.033 | 0.000 | 8,338.033 | 15,392.967 | 23,731.000 | 9,247.637 |
| HI_CEM | 31,570.704 | 8,401.633 | 39,972.337 | 875,170.105 | 빈칸 | 빈칸 | 5,712,566.000 | 370,165.000 | 6,082,731.000 | 빈칸 |
| HI_SEC | 1,086.296 | 917.367 | 2,003.663 | 95,559.063 | 빈칸 | 빈칸 | 54,947.000 | 40,418.000 | 95,365.000 | 빈칸 |
| KG 합계 | 2,199.462 | 6,485.538 | 8,685.000 | 675,576.877 | | | 111,253.115 | 291,283.885 | 402,537.000 | |
| TW 합계 | 755.873 | 1,997.127 | 2,753.000 | 208,034.064 | | | 52,690.343 | 98,580.657 | 151,271.000 | |
| HBG 합계 | 251.091 | 1,168.909 | 1,420.000 | 121,761.354 | | | 12,545.005 | 52,325.995 | 64,871.000 | |
| SA 합계 | 164.842 | 310.158 | 475.000 | 32,308.127 | | | 8,338.033 | 15,392.967 | 23,731.000 | |
| HI 합계 | 32,657.000 | 9,319.000 | 41,976.000 | 970,729.168 | | | 5,767,513.000 | 410,583.000 | 6,178,096.000 | |
"""

# 표 2. 월 값 — 에너지·배출
TABLE_2 = """
| plant_id | 칸 | 1월 | 2월 | 3월 | 4월 | 5월 | 6월 | 7월 | 8월 | 9월 | 10월 | 11월 | 12월 |
|---|---|---|---|---|---|---|---|---|---|---|---|---|---|
| KG_DJ | fuel_tj_gross | 152.330 | 117.353 | 127.483 | 138.760 | 133.599 | 133.599 | 143.729 | 143.538 | 143.538 | 133.217 | 108.944 | 124.425 |
| KG_DJ | elec_tj | 397.690 | 374.931 | 400.679 | 388.298 | 384.005 | 396.089 | 423.342 | 409.245 | 410.810 | 372.572 | 365.562 | 396.203 |
| KG_DJ | direct_t | 7,705.144 | 5,935.958 | 6,448.345 | 7,018.738 | 6,757.711 | 6,757.711 | 7,270.098 | 7,260.430 | 7,260.430 | 6,738.376 | 5,510.580 | 6,293.662 |
| KG_DJ | indirect_t | 17,861.382 | 16,839.206 | 17,995.628 | 17,439.561 | 17,246.745 | 17,789.506 | 19,013.474 | 18,380.339 | 18,450.653 | 16,733.271 | 16,418.420 | 17,794.624 |
| KG_IC | fuel_tj_gross | 57.005 | 43.916 | 47.707 | 51.927 | 49.996 | 49.996 | 53.786 | 53.715 | 53.715 | 49.853 | 40.769 | 46.562 |
| KG_IC | elec_tj | 150.289 | 145.719 | 150.362 | 138.674 | 144.239 | 146.741 | 154.243 | 149.563 | 155.362 | 135.112 | 144.538 | 151.270 |
| KG_IC | direct_t | 2,883.432 | 2,221.364 | 2,413.111 | 2,626.564 | 2,528.882 | 2,528.882 | 2,720.628 | 2,717.010 | 2,717.010 | 2,521.646 | 2,062.178 | 2,355.225 |
| KG_IC | indirect_t | 6,749.886 | 6,544.662 | 6,753.168 | 6,228.260 | 6,478.199 | 6,590.556 | 6,927.482 | 6,717.290 | 6,977.742 | 6,068.272 | 6,491.616 | 6,793.943 |
| TW_HJ | fuel_tj_gross | 18.909 | 17.344 | 19.052 | 19.377 | 19.458 | 19.174 | 19.539 | 18.665 | 19.295 | 17.933 | 17.933 | 18.787 |
| TW_HJ | elec_tj | 80.081 | 77.636 | 80.553 | 80.624 | 75.150 | 76.658 | 83.508 | 80.726 | 85.428 | 70.706 | 74.985 | 79.411 |
| TW_HJ | direct_fuel_t | 956.461 | 877.270 | 963.660 | 980.116 | 984.229 | 969.831 | 988.343 | 944.120 | 976.002 | 907.096 | 907.095 | 950.291 |
| TW_HJ | process_t | 1,212.447 | 1,112.062 | 1,221.573 | 1,242.432 | 1,247.647 | 1,229.395 | 1,252.862 | 1,196.802 | 1,237.217 | 1,149.869 | 1,149.869 | 1,204.625 |
| TW_HJ | direct_t | 2,168.908 | 1,989.332 | 2,185.233 | 2,222.548 | 2,231.876 | 2,199.226 | 2,241.205 | 2,140.922 | 2,213.219 | 2,056.965 | 2,056.964 | 2,154.916 |
| TW_HJ | indirect_t | 3,952.877 | 3,832.226 | 3,976.197 | 3,979.670 | 3,709.493 | 3,783.951 | 4,122.046 | 3,984.730 | 4,216.832 | 3,490.143 | 3,701.366 | 3,919.839 |
| TW_NS | fuel_tj_gross | 44.484 | 40.800 | 44.818 | 45.584 | 45.775 | 45.105 | 45.966 | 43.910 | 45.392 | 42.188 | 42.188 | 44.197 |
| TW_NS | elec_tj | 89.075 | 86.356 | 89.601 | 89.679 | 83.591 | 85.269 | 92.887 | 89.793 | 95.023 | 78.648 | 83.408 | 88.331 |
| TW_NS | direct_t | 2,250.067 | 2,063.771 | 2,267.003 | 2,305.714 | 2,315.392 | 2,281.520 | 2,325.070 | 2,221.034 | 2,296.037 | 2,133.935 | 2,133.935 | 2,235.551 |
| TW_NS | indirect_t | 4,396.865 | 4,262.663 | 4,422.804 | 4,426.667 | 4,126.144 | 4,208.966 | 4,585.035 | 4,432.296 | 4,690.468 | 3,882.158 | 4,117.105 | 4,360.116 |
| HBG_CW | fuel_tj_gross | 23.633 | 18.207 | 19.779 | 21.528 | 20.727 | 20.727 | 22.299 | 22.269 | 22.269 | 20.668 | 16.902 | 19.304 |
| HBG_CW | elec_tj | 96.292 | 96.788 | 98.689 | 97.984 | 93.249 | 95.275 | 102.410 | 94.330 | 103.368 | 85.621 | 93.247 | 98.720 |
| HBG_CW | direct_t | 1,180.764 | 909.648 | 988.168 | 1,075.577 | 1,035.576 | 1,035.576 | 1,114.096 | 1,112.615 | 1,112.614 | 1,032.613 | 844.461 | 964.463 |
| HBG_CW | indirect_t | 4,310.475 | 4,332.708 | 4,417.774 | 4,386.232 | 4,174.294 | 4,264.954 | 4,584.351 | 4,222.678 | 4,627.262 | 3,832.809 | 4,174.200 | 4,419.171 |
| HBG_GJ | fuel_tj_gross | 0.221 | 0.170 | 0.185 | 0.201 | 0.194 | 0.194 | 0.208 | 0.208 | 0.208 | 0.193 | 0.158 | 0.181 |
| HBG_GJ | elec_tj | 0.932 | 0.925 | 0.933 | 0.930 | 0.852 | 0.867 | 0.973 | 0.927 | 0.962 | 0.794 | 0.824 | 0.884 |
| HBG_GJ | direct_t | 11.035 | 8.501 | 9.235 | 10.052 | 9.678 | 9.678 | 10.412 | 10.398 | 10.398 | 9.651 | 7.892 | 9.014 |
| HBG_GJ | indirect_t | 41.732 | 41.426 | 41.786 | 41.642 | 38.146 | 38.820 | 43.537 | 41.478 | 43.052 | 35.526 | 36.904 | 39.561 |
| HBG_SA | fuel_tj_gross | 0.044 | 0.034 | 0.036 | 0.040 | 0.038 | 0.038 | 0.041 | 0.041 | 0.041 | 0.038 | 0.031 | 0.036 |
| HBG_SA | elec_tj | 0.180 | 0.169 | 0.181 | 0.176 | 0.174 | 0.179 | 0.191 | 0.185 | 0.186 | 0.168 | 0.165 | 0.179 |
| HBG_SA | direct_t | 2.179 | 1.678 | 1.823 | 1.984 | 1.911 | 1.911 | 2.056 | 2.053 | 2.053 | 1.905 | 1.558 | 1.779 |
| HBG_SA | indirect_t | 8.046 | 7.585 | 8.106 | 7.856 | 7.769 | 8.013 | 8.564 | 8.279 | 8.311 | 7.537 | 7.396 | 8.015 |
| SA_PS | fuel_tj_gross | 13.182 | 13.262 | 13.279 | 14.210 | 14.645 | 14.538 | 15.316 | 14.004 | 14.845 | 12.634 | 11.970 | 12.957 |
| SA_PS | elec_tj | 26.423 | 24.579 | 26.475 | 24.783 | 24.174 | 25.250 | 27.717 | 27.702 | 26.974 | 24.782 | 25.028 | 26.271 |
| SA_PS | direct_t | 666.751 | 670.826 | 671.674 | 718.790 | 740.777 | 735.344 | 774.735 | 708.348 | 750.880 | 639.075 | 605.458 | 655.375 |
| SA_PS | indirect_t | 1,311.345 | 1,219.847 | 1,313.965 | 1,229.982 | 1,199.730 | 1,253.138 | 1,375.579 | 1,374.837 | 1,338.711 | 1,229.895 | 1,242.132 | 1,303.806 |
| HI_CEM | fuel_tj_gross | 2,047.750 | 2,051.433 | 2,791.716 | 3,204.213 | 2,942.720 | 2,843.279 | 2,666.494 | 2,353.439 | 2,677.543 | 2,335.024 | 2,883.792 | 2,773.301 |
| HI_CEM | elec_tj | 675.085 | 641.407 | 757.103 | 691.137 | 680.413 | 699.532 | 732.073 | 717.043 | 729.951 | 682.415 | 678.234 | 717.240 |
| HI_CEM | direct_t | 370,530.412 | 371,196.834 | 505,147.577 | 579,786.796 | 532,470.863 | 514,477.479 | 482,489.242 | 425,843.406 | 484,488.507 | 422,511.298 | 521,808.117 | 501,815.469 |
| HI_CEM | indirect_t | 29,743.347 | 28,259.549 | 33,356.965 | 30,450.598 | 29,978.127 | 30,820.458 | 32,254.204 | 31,591.967 | 32,160.673 | 30,066.319 | 29,882.113 | 31,600.680 |
| HI_SEC | fuel_tj_gross | 70.460 | 70.586 | 96.058 | 110.252 | 101.254 | 97.833 | 91.750 | 80.978 | 92.130 | 80.344 | 99.226 | 95.425 |
| HI_SEC | elec_tj | 78.378 | 73.033 | 78.153 | 76.072 | 75.275 | 75.731 | 80.882 | 79.352 | 79.844 | 72.143 | 72.391 | 76.113 |
| HI_SEC | direct_t | 3,563.991 | 3,570.401 | 4,858.822 | 5,576.749 | 5,121.635 | 4,948.563 | 4,640.881 | 4,096.026 | 4,660.111 | 4,063.975 | 5,019.074 | 4,826.772 |
| HI_SEC | indirect_t | 3,453.228 | 3,217.761 | 3,443.312 | 3,351.629 | 3,316.512 | 3,336.600 | 3,563.564 | 3,496.140 | 3,517.818 | 3,178.549 | 3,189.449 | 3,353.438 |
"""

# 표 3. 월 값 — 생산 (quantity). 줄 순서 = 생산 파일 line_no 1~12
TABLE_3 = """
| plant_id | product | 단위 | 1월 | 2월 | 3월 | 4월 | 5월 | 6월 | 7월 | 8월 | 9월 | 10월 | 11월 | 12월 |
|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|
| KG_DJ | 판재 | t | 152,851.923 | 117,755.434 | 127,919.991 | 139,235.252 | 134,057.082 | 134,057.081 | 144,221.638 | 144,029.854 | 144,029.854 | 133,673.513 | 109,316.933 | 124,851.445 |
| KG_IC | 판재 | t | 57,200.502 | 44,066.635 | 47,870.432 | 52,104.848 | 50,167.065 | 50,167.065 | 53,970.862 | 53,899.093 | 53,899.092 | 50,023.525 | 40,908.765 | 46,722.116 |
| TW_HJ | 잉곳·블룸 | t | 15,155.586 | 13,900.769 | 15,269.661 | 15,530.402 | 15,595.587 | 15,367.439 | 15,660.773 | 14,960.031 | 15,465.217 | 14,373.363 | 14,373.363 | 15,057.809 |
| TW_NS | 단조품 | t | 5,716.787 | 5,243.462 | 5,759.817 | 5,858.170 | 5,882.758 | 5,796.699 | 5,907.346 | 5,643.022 | 5,833.581 | 5,421.727 | 5,421.727 | 5,679.904 |
| HBG_CW | 스테인리스 강판 | t | 22,788.071 | 17,555.678 | 19,071.070 | 20,758.017 | 19,986.024 | 19,986.024 | 21,501.417 | 21,472.825 | 21,472.824 | 19,928.840 | 16,297.616 | 18,613.594 |
| HBG_GJ | 자동차부품 등 | 개 | 13,531 | 10,424 | 11,324 | 12,326 | 11,867 | 11,867 | 12,767 | 12,750 | 12,750 | 11,834 | 9,677 | 11,053 |
| HBG_SA | 스테인리스 표면 가공품 | t | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 |
| SA_PS | 압연 | t | 1,236.019 | 1,243.573 | 1,245.146 | 1,332.489 | 1,373.249 | 1,363.177 | 1,436.199 | 1,313.132 | 1,391.977 | 1,184.715 | 1,122.394 | 1,214.930 |
| SA_PS | 가공 | km | 9,512.234 | 9,570.368 | 9,582.479 | 10,254.658 | 10,568.341 | 10,490.829 | 11,052.794 | 10,105.689 | 10,712.466 | 9,117.404 | 8,637.796 | 9,349.942 |
| HI_CEM | 시멘트 외 | t | 600,625.292 | 601,705.553 | 818,838.077 | 939,827.345 | 863,128.791 | 833,961.736 | 782,109.193 | 690,286.981 | 785,349.977 | 684,885.674 | 845,844.610 | 813,436.771 |
| HI_SEC | 레미콘 | ㎥ | 131,151.656 | 131,387.541 | 178,800.280 | 205,219.319 | 188,471.535 | 182,102.660 | 170,780.215 | 150,730.051 | 171,487.867 | 149,550.630 | 184,697.387 | 177,620.859 |
| HI_SEC | 레미탈 | t | 208,986.468 | 209,362.342 | 284,913.206 | 327,011.199 | 300,324.078 | 290,175.455 | 272,133.458 | 240,184.088 | 273,261.083 | 238,304.713 | 294,310.079 | 283,033.831 |
"""

# 표 4. 월 값 — 원재료 매입 (quantity_t)
TABLE_4 = """
| line_no | plant_id | material | 1월 | 2월 | 3월 | 4월 | 5월 | 6월 | 7월 | 8월 | 9월 | 10월 | 11월 | 12월 | 연간 |
|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|
| 1 | KG_DJ | 열연 코일 | 141,462.134 | 108,980.866 | 118,388.010 | 128,860.112 | 124,067.794 | 124,067.794 | 133,474.937 | 133,297.444 | 133,297.444 | 123,712.808 | 101,171.163 | 115,548.117 | 1,486,328.623 |
| 2 | KG_DJ | 기타 코일 | 16,830.610 | 12,966.116 | 14,085.341 | 15,331.271 | 14,761.100 | 14,761.100 | 15,880.325 | 15,859.208 | 15,859.208 | 14,718.865 | 12,036.949 | 13,747.462 | 176,837.555 |
| 3 | KG_IC | 열연 코일 | 52,938.196 | 40,783.002 | 44,303.358 | 48,222.246 | 46,428.857 | 46,428.857 | 49,949.214 | 49,882.792 | 49,882.792 | 46,296.013 | 37,860.441 | 43,240.609 | 556,216.377 |
| 4 | KG_IC | 기타 코일 | 6,298.379 | 4,852.202 | 5,271.040 | 5,737.294 | 5,523.924 | 5,523.923 | 5,942.762 | 5,934.859 | 5,934.859 | 5,508.118 | 4,504.487 | 5,144.598 | 66,176.445 |
| 5 | HBG_CW | 스테인리스 열연 강대 | 22,788.071 | 17,555.678 | 19,071.070 | 20,758.017 | 19,986.024 | 19,986.024 | 21,501.417 | 21,472.825 | 21,472.824 | 19,928.840 | 16,297.616 | 18,613.594 | 239,432.000 |
| 6 | HBG_SA | 창원 판 (이동분) | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 |
| 7 | TW_HJ | 철 스크랩 | 17,521.975 | 16,071.231 | 17,653.861 | 17,955.314 | 18,030.677 | 17,766.906 | 18,106.040 | 17,295.885 | 17,879.951 | 16,617.615 | 16,617.615 | 17,408.930 | 208,926.000 |
| 8 | TW_HJ | 철스크랩 외의 "외" | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 |
| 9 | TW_NS | 화전 잉곳·블룸 (이동분) | 5,716.787 | 5,243.462 | 5,759.817 | 5,858.170 | 5,882.758 | 5,796.699 | 5,907.346 | 5,643.022 | 5,833.581 | 5,421.727 | 5,421.727 | 5,679.904 | 68,165.000 |
| 10 | TW_NS | 단조 원재료 외부 매입분 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 |
| 11 | SA_PS | AL-STRIP | 1,236.019 | 1,243.573 | 1,245.146 | 1,332.489 | 1,373.249 | 1,363.177 | 1,436.199 | 1,313.132 | 1,391.977 | 1,184.715 | 1,122.394 | 1,214.930 | 15,457.000 |
| 12 | HI_CEM | 클링커 매입 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 | 0.000 |
"""

# 표 5. EU 신고 표 숫자 칸
TABLE_5 = """
| row_no | plant_id | report_case | gov_total_t | own_t | precursor_t | embedded_t | production_t | see_t_per_t | precursor_x110_2026_t | embedded_markup_2026_t |
|---|---|---|---|---|---|---|---|---|---|---|
| 1 | KG_DJ | 원산지 증빙 없을 때 | 292,919.992 | 80,957.183 | 13,691,394.396 | 13,772,351.579 | 1,606,000.000 | 8.5756 | 15,060,533.836 | 15,141,491.019 |
| 2 | KG_DJ | 한국산 증빙 시 | 292,919.992 | 80,957.183 | 3,527,183.741 | 3,608,140.924 | 1,606,000.000 | 2.2467 | 3,879,902.115 | 3,960,859.298 |
| 3 | KG_IC | 원산지 증빙 없을 때 | 109,617.008 | 30,295.932 | 5,123,616.454 | 5,153,912.386 | 601,000.000 | 8.5756 | 5,635,978.099 | 5,666,274.031 |
| 4 | KG_IC | 한국산 증빙 시 | 109,617.008 | 30,295.932 | 1,319,948.585 | 1,350,244.517 | 601,000.000 | 2.2467 | 1,451,943.444 | 1,482,239.376 |
| 5 | TW_HJ | 해당 없음 (원재료 없음) | 72,530.684 | 25,861.314 | 0.000 | 25,861.314 | 180,710.000 | 0.1431 | 빈칸 | 빈칸 |
| 6 | TW_NS | 이동분 | 78,740.316 | 26,829.029 | 9,755.058 | 36,584.087 | 68,165.000 | 0.5367 | 빈칸 | 빈칸 |
| 7 | HBG_CW | 원산지 증빙 없을 때 | 64,153.079 | 12,406.171 | 2,075,875.440 | 2,088,281.611 | 239,432.000 | 8.7218 | 2,283,462.984 | 2,295,869.155 |
| 8 | HBG_CW | 한국산 증빙 시 | 64,153.079 | 12,406.171 | 869,138.160 | 881,544.331 | 239,432.000 | 3.6818 | 956,051.976 | 968,458.147 |
| 9 | HBG_SA | 이동 수량 미공시 | 118.367 | 0.000 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 |
| 10 | HBG_GJ | 비해당 | 599.554 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 |
| 11 | SA_PS | 원산지 증빙 없을 때 | 23,731.000 | 8,338.033 | 77,068.602 | 85,406.635 | 15,457.000 | 5.5254 | 84,775.462 | 93,113.495 |
| 12 | SA_PS | 한국산 증빙 시 | 23,731.000 | 8,338.033 | 30,141.150 | 38,479.183 | 15,457.000 | 2.4894 | 33,155.265 | 41,493.298 |
| 13 | HI_CEM | 해당 없음 (원재료 없음) | 6,082,731.000 | 6,082,731.000 | 0.000 | 6,082,731.000 | 9,260,000.000 | 0.6569 | 빈칸 | 빈칸 |
| 14 | HI_SEC | 비해당 | 95,365.000 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 | 빈칸 |
"""

# 표 6. 은행 전달 표 숫자 칸
TABLE_6 = """
| row_no | plant_id | direct_t | indirect_t | total_t | fuel_tj | elec_mwh | elec_factor_check | elec_factor_diff_pct | precursor_unknown_t | precursor_kr_t | eu_embedded_unknown_t | eu_embedded_kr_t |
|---|---|---|---|---|---|---|---|---|---|---|---|---|
| 1 | KG_ALL | 111,253.115 | 291,283.885 | 402,537.000 | 2,199.462 | 675,576.877 | 0.4312 | -6.1 | 18,815,010.850 | 4,847,132.326 | 18,926,263.965 | 4,958,385.441 |
| 2 | KG_DJ | 80,957.183 | 211,962.809 | 292,919.992 | 1,600.515 | 491,606.876 | 0.4312 | -6.1 | 13,691,394.396 | 3,527,183.741 | 13,772,351.579 | 3,608,140.924 |
| 3 | KG_IC | 30,295.932 | 79,321.076 | 109,617.008 | 598.947 | 183,970.001 | 0.4312 | -6.1 | 5,123,616.454 | 1,319,948.585 | 5,153,912.386 | 1,350,244.517 |
| 4 | HI_ALL | 5,767,513.000 | 410,583.000 | 6,178,096.000 | 32,657.000 | 970,729.168 | 0.4230 | -7.9 | 0.000 | 0.000 | 6,082,731.000 | 6,082,731.000 |
| 5 | HI_CEM | 5,712,566.000 | 370,165.000 | 6,082,731.000 | 31,570.704 | 875,170.105 | 0.4230 | -7.9 | 0.000 | 0.000 | 6,082,731.000 | 6,082,731.000 |
| 6 | HI_SEC | 54,947.000 | 40,418.000 | 95,365.000 | 1,086.296 | 95,559.063 | 0.4230 | -7.9 | 0.000 | 0.000 | 빈칸 | 빈칸 |
| 7 | SA_ALL | 8,338.033 | 15,392.967 | 23,731.000 | 164.842 | 32,308.127 | 0.4764 | 3.7 | 77,068.602 | 30,141.150 | 85,406.635 | 38,479.183 |
| 8 | SA_PS | 8,338.033 | 15,392.967 | 23,731.000 | 164.842 | 32,308.127 | 0.4764 | 3.7 | 77,068.602 | 30,141.150 | 85,406.635 | 38,479.183 |
"""

# #14 문구 (09 4-8-6·4-8-8 표 문구를 다시 적음)
PLANT_MASTER_TEXT = {
    "KG_DJ": ("해당 — 석도 라인(plating 해석) 배출 포함", "7209·7210"),
    "KG_IC": ("해당 — 석도 라인(plating 해석) 배출 포함", "7209·7210"),
    "TW_HJ": ("해당 — 잉곳·블룸 (대부분 녹산으로 이동)", "7206·7207·7218 등 (품목 추정)"),
    "TW_NS": ("일부 해당 — 비율 미공시",
              "7308(타워 플랜지 — 대법원 2011두12290) / 7326 또는 84류(그 외 — 불명)"),
    "TW_DD": ("숫자 줄 없음 (사업부 단위로 나눔)", "불명"),
    "HBG_CW": ("해당", "7219·7220"),
    "HBG_GJ": ("비해당 — 자동차 엔진 블럭용 사형코어", ""),
    "HBG_SA": ("EU 대상 물건(추정) — 자체 공정 배출 제외(마무리)", "7219·7220 (추정)"),
    "SA_PS": ("해당 — 가공(라미네이팅·인쇄, finishing 해석) 배출 포함", "7607"),
    "HI_CEM": ("해당 — 단양·삼곡·영월 포틀랜드, 당진·평택 일부 해당(슬래그시멘트 해당, 고로슬래그 미분말 비해당) — 비대상 미분말 몫 섞임",
               "2523"),
    "HI_SEC": ("비해당 — 레미콘·레미탈 (2026/1740 한국 표에 없음)", ""),
}
_UNKNOWN = ("원산지 증빙 없을 때", "원산지 불명 최고값 — 특정 공장 값 아님")
_KOREA = ("한국산 증빙 시", "국가 평균값 — 특정 공장 값 아님")
EU_REPORT_TEXT = {
    "1": _UNKNOWN, "2": _KOREA, "3": _UNKNOWN, "4": _KOREA,
    "5": ("해당 없음 (원재료 없음)", "합성 실제값 (자체 공정)"),
    "6": ("이동분", "회사 안 이동 — 보내는 공장 합성 값 (EU 기본값 아님)"),
    "7": _UNKNOWN, "8": _KOREA,
    "9": ("이동 수량 미공시", "이동 수량 미공시"),
    "10": ("비해당", "비해당 — 정부 신고에만 있음"),
    "11": _UNKNOWN, "12": _KOREA,
    "13": ("해당 없음 (원재료 없음)", "회사 사업별 공개값 (자체 공정)"),
    "14": ("비해당", "비해당 — 정부 신고에만 있음"),
}
_BANK_TOTAL = "합계 = 실측(명세서 2025), 직접·간접 = 합성 — 실제로는 회사 실측 자리"
_BANK_PLANT = "합성 — 실제로는 회사 실측 자리"
_BANK_HI_GROUP = "회사 공개값 (배출) · 합성 (에너지 — 결정 39)"
BANK_TEXT = {"1": _BANK_TOTAL, "2": _BANK_PLANT, "3": _BANK_PLANT,
             "4": "실측 (명세서 2025 = 지속가능경영보고서)", "5": _BANK_HI_GROUP,
             "6": _BANK_HI_GROUP, "7": _BANK_TOTAL, "8": _BANK_PLANT}
ENVIZI_TEXT = "Envizi 입력값: elec_mwh, fuel_tj_gross"

# #15·#16·#17 파일 (09 4-8-3 순서) — (경로, 행 수, 크기 상한 KB, 원천)
FILES = [
    ("data/demo/mfg/ext/mfg_ext_anchor.csv", 5, 5, "③"),
    ("data/demo/mfg/ext/mfg_ext_factors.xlsx", None, 30, "③"),
    ("data/demo/mfg/ext/mfg_ext_stats.csv", 171, 60, "③"),
    ("data/demo/mfg/ext/mfg_ext_disclosed.csv", 53, 30, "③"),
    ("data/demo/mfg/db/mfg_plant_master_synth.csv", 11, 15, "①"),
    ("data/demo/mfg/db/mfg_energy_monthly_synth.csv", 120, 120, "①"),
    ("data/demo/mfg/db/mfg_production_monthly_synth.csv", 144, 100, "①"),
    ("data/demo/mfg/db/mfg_emissions_monthly_synth.csv", 120, 120, "①"),
    ("data/demo/mfg/db/mfg_purchases_monthly_synth.csv", 144, 160, "①"),
    ("data/demo/mfg/db/mfg_eu_report_synth.csv", 14, 20, "①"),
    ("data/demo/mfg/db/mfg_bank_transfer_synth.csv", 8, 20, "①"),
]
SHEET_ROWS = {"eu_defaults": 12, "gov_factors": 6, "markup": 5}
KB = 1000  # 1 KB = 1,000 B (사용자 확정 2026-10-12 — 더 엄격한 쪽)

# ── 읽기·비교 ─────────────────────────────────────────────────────────────


def _read(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _d(text):
    """문자 → Decimal, 빈 문자열 → None"""
    return None if text in ("", None) else Decimal(text)


def _q(x, n):
    """소수 n자리 반올림 (ROUND_HALF_UP — 0에서 먼 쪽)"""
    return x.quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP)


def _md(text):
    """마크다운 표 → (머리글, 값 줄 목록). 칸: `빈칸` → None, 비어 있음 → ''(비교하지 않음), 그 밖 → Decimal 또는 문자"""
    lines = [ln.strip() for ln in text.strip().splitlines()]
    head = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = []
    for ln in lines[2:]:
        cells = [c.strip() for c in ln.strip("|").split("|")]
        rows.append(dict(zip(head, cells)))
    return head, rows


def _exp(cell):
    """표의 숫자 칸 → 기대값 (Decimal / None = 빈칸 / '' = 표에 없음)"""
    if cell == "":
        return ""
    if cell == "빈칸":
        return None
    return Decimal(cell.replace(",", ""))


def _sum(values):
    """합 — 모두 빈칸이면 None, 일부만 빈칸이면 '일부 빈칸'"""
    vals = [v for v in values if v is not None]
    if not vals:
        return None
    if len(vals) != len(values):
        return "일부 빈칸"
    return sum(vals, Decimal(0))


def _same(expected, actual):
    if expected is None or actual is None or isinstance(actual, str):
        return expected is None and actual is None
    return expected == actual


def _s(v):
    return "빈칸" if v is None else str(v)


def _miss(expected, actual, *ctx):
    where = ", ".join(str(c) for c in ctx if c not in (None, ""))
    return f"기대 {_s(expected)} 실제 {_s(actual)}" + (f" ({where})" if where else "")


def _load():
    d = {"energy": _read(DB / "mfg_energy_monthly_synth.csv"),
         "emissions": _read(DB / "mfg_emissions_monthly_synth.csv"),
         "production": _read(DB / "mfg_production_monthly_synth.csv"),
         "purchases": _read(DB / "mfg_purchases_monthly_synth.csv"),
         "plant_master": _read(DB / "mfg_plant_master_synth.csv"),
         "eu_report": _read(DB / "mfg_eu_report_synth.csv"),
         "bank": _read(DB / "mfg_bank_transfer_synth.csv")}
    d["en"] = {(r["plant_id"], r["month"]): r for r in d["energy"]}
    d["em"] = {(r["plant_id"], r["month"]): r for r in d["emissions"]}
    d["pr"] = {(r["line_no"], r["month"]): r for r in d["production"]}
    d["pu"] = {(r["line_no"], r["month"]): r for r in d["purchases"]}
    d["eu"] = {r["row_no"]: r for r in d["eu_report"]}
    d["bk"] = {r["row_no"]: r for r in d["bank"]}
    return d


def _company_sum(d, file, col):
    out = {}
    for r in d[file]:
        out[r["company_id"]] = out.get(r["company_id"], Decimal(0)) + Decimal(r[col])
    return out


ENERGY_COLS = {"fuel_tj_gross", "fuel_tj_net", "elec_tj", "elec_mwh", "energy_tj"}


def _cell(d, pid, month, col):
    src = d["en"] if col in ENERGY_COLS else d["em"]
    return _d(src[(pid, month)][col])


# ── 검산 17개 (함수 docstring 첫 줄 = 4-8-10 "검산" 칸 문구) ──────────────────


def check_01(d):
    """검산 #01 — 회사별 Σ total_t = 앵커 배출"""
    got = _company_sum(d, "emissions", "total_t")
    return [_miss(v, got.get(c), c, "total_t") for c, v in ANCHOR_EMISSIONS.items()
            if got.get(c) != Decimal(v)]


def check_02(d):
    """검산 #02 — 회사별 Σ energy_tj = 앵커 에너지"""
    got = _company_sum(d, "energy", "energy_tj")
    return [_miss(v, got.get(c), c, "energy_tj") for c, v in ANCHOR_ENERGY.items()
            if got.get(c) != Decimal(v)]


def check_03(d):
    """검산 #03 — 행마다 total_t = direct_t + indirect_t, energy_tj = fuel_tj_gross + elec_tj, KG·TW·SA 행은 direct_t = direct_fuel_t + process_t"""
    fails = []
    for r in d["emissions"]:
        direct, indirect, total = _d(r["direct_t"]), _d(r["indirect_t"]), _d(r["total_t"])
        if total != direct + indirect:
            fails.append(_miss(direct + indirect, total, r["plant_id"], "total_t", r["month"]))
        if r["company_id"] in FUEL_FIRST_COMPANIES:
            s = _d(r["direct_fuel_t"]) + _d(r["process_t"])
            if direct != s:
                fails.append(_miss(s, direct, r["plant_id"], "direct_t", r["month"]))
    for r in d["energy"]:
        s = _d(r["fuel_tj_gross"]) + _d(r["elec_tj"])
        if _d(r["energy_tj"]) != s:
            fails.append(_miss(s, r["energy_tj"], r["plant_id"], "energy_tj", r["month"]))
    return fails


def check_04(d):
    """검산 #04 — 공장 연간 값"""
    alias = {"Σ월 elec_mwh": "elec_mwh", "Σ월 direct_fuel_alt_net_t": "direct_fuel_alt_net_t"}
    head, rows = _md(TABLE_1)
    fails = []
    for row in rows:
        label = row["plant_id"]
        plants = COMPANY_PLANTS[label.split()[0]] if label.endswith(" 합계") else [label]
        for h in head[1:]:
            exp = _exp(row[h])
            if exp == "":
                continue
            col = alias.get(h, h)
            got = _sum([_cell(d, p, m, col) for p in plants for m in MONTHS])
            if not _same(exp, got):
                fails.append(_miss(exp, got, label, col))
    return fails


def check_05(d):
    """검산 #05 — 월 값"""
    fails = []
    _, rows = _md(TABLE_2)
    for row in rows:
        pid, col = row["plant_id"], row["칸"]
        for k, m in enumerate(MONTHS):
            exp, got = _exp(row[f"{k + 1}월"]), _cell(d, pid, m, col)
            if not _same(exp, got):
                fails.append(_miss(exp, got, pid, col, m))
    # 표 2 머리 설명: KG·SA·TW_NS는 direct_fuel_t = direct_t, process_t = 0.000
    for pid in DIRECT_EQ_FUEL:
        for m in MONTHS:
            direct = _cell(d, pid, m, "direct_t")
            fuel, proc = _cell(d, pid, m, "direct_fuel_t"), _cell(d, pid, m, "process_t")
            if not _same(direct, fuel):
                fails.append(_miss(direct, fuel, pid, "direct_fuel_t", m))
            if not _same(Decimal("0.000"), proc):
                fails.append(_miss("0.000", proc, pid, "process_t", m))
    # 4-8-6 배출 파일: HBG·HI는 direct_fuel_t·process_t 빈칸
    for pid in DIRECT_SPLIT_BLANK:
        for m in MONTHS:
            for col in ("direct_fuel_t", "process_t"):
                got = _cell(d, pid, m, col)
                if got is not None:
                    fails.append(_miss(None, got, pid, col, m))
    # 표 3: 줄 순서 = line_no
    _, rows = _md(TABLE_3)
    for i, row in enumerate(rows, start=1):
        line = str(i)
        for k, m in enumerate(MONTHS):
            r = d["pr"][(line, m)]
            if r["plant_id"] != row["plant_id"]:
                fails.append(_miss(row["plant_id"], r["plant_id"], f"{line}번 줄", "plant_id", m))
            exp, got = _exp(row[f"{k + 1}월"]), _d(r["quantity"])
            if not _same(exp, got):
                fails.append(_miss(exp, got, row["plant_id"], f"quantity {line}번 줄", m))
    # 표 4: 12달 값과 연간 = 12달 합 (6번 줄은 12달 모두 빈칸인지만)
    _, rows = _md(TABLE_4)
    for row in rows:
        line = row["line_no"]
        got_months = []
        for k, m in enumerate(MONTHS):
            r = d["pu"][(line, m)]
            exp, got = _exp(row[f"{k + 1}월"]), _d(r["quantity_t"])
            got_months.append(got)
            if not _same(exp, got):
                fails.append(_miss(exp, got, row["plant_id"], f"quantity_t {line}번 줄", m))
        exp = _exp(row["연간"])
        if exp is None:
            if any(g is not None for g in got_months):
                fails.append(_miss(None, "값 있음", row["plant_id"], f"quantity_t {line}번 줄 12달"))
        else:
            got = _sum(got_months)
            if not _same(exp, got):
                fails.append(_miss(exp, got, row["plant_id"], f"quantity_t {line}번 줄 연간"))
    return fails


def check_06(d):
    """검산 #06 — 행마다 elec_mwh = r3(elec_tj × 1000 ÷ 9.6) (120행). KG·TW·SA 공장 60행은 fuel_tj_net = r3(fuel_tj_gross × NET), direct_fuel_alt_net_t = r3(fuel_tj_gross × GAS). HBG·HI 공장 60행은 두 칸 빈칸"""
    fails = []
    for r in d["energy"]:
        pid, m = r["plant_id"], r["month"]
        exp = _q(_d(r["elec_tj"]) * 1000 / MJ, 3)
        if _d(r["elec_mwh"]) != exp:
            fails.append(_miss(exp, r["elec_mwh"], pid, "elec_mwh", m))
        fuel = _d(r["fuel_tj_gross"])
        net, alt = _d(r["fuel_tj_net"]), _d(d["em"][(pid, m)]["direct_fuel_alt_net_t"])
        if r["company_id"] in FUEL_FIRST_COMPANIES:
            exp_net, exp_alt = _q(fuel * NET, 3), _q(fuel * GAS, 3)
        else:
            exp_net = exp_alt = None
        if not _same(exp_net, net):
            fails.append(_miss(exp_net, net, pid, "fuel_tj_net", m))
        if not _same(exp_alt, alt):
            fails.append(_miss(exp_alt, alt, pid, "direct_fuel_alt_net_t", m))
    return fails


def check_07(d):
    """검산 #07 — 공장마다 Σ월 elec_mwh 대 r3(공장 연간 elec_tj × 1000 ÷ 9.6)"""
    fails = []
    for pid in [p for ps in COMPANY_PLANTS.values() for p in ps]:
        mwh = _sum([_cell(d, pid, m, "elec_mwh") for m in MONTHS])
        annual = _q(_sum([_cell(d, pid, m, "elec_tj") for m in MONTHS]) * 1000 / MJ, 3)
        if abs(mwh - annual) > TOL_07:
            fails.append(_miss(annual, mwh, pid, f"Σ월 elec_mwh — 차이 {abs(mwh - annual)} > {TOL_07}"))
    return fails


def check_08(d):
    """검산 #08 — KG·TW·SA 행마다 direct_fuel_t와 fuel_tj_gross × NET × GAS의 차이(절댓값)"""
    fails = []
    for r in d["emissions"]:
        if r["company_id"] not in FUEL_FIRST_COMPANIES:
            continue
        pid, m = r["plant_id"], r["month"]
        calc = _d(d["en"][(pid, m)]["fuel_tj_gross"]) * NET * GAS
        diff = abs(_d(r["direct_fuel_t"]) - calc)
        if diff > TOL_08:
            fails.append(_miss(_q(calc, 3), r["direct_fuel_t"], pid,
                               f"direct_fuel_t — 차이 {_q(diff, 4)} > {TOL_08}", m))
    return fails


def check_09(d):
    """검산 #09 — 대조: 회사 Σ indirect_t ÷ Σ elec_mwh (넷째 자리)"""
    ind = _company_sum(d, "emissions", "indirect_t")
    mwh = _company_sum(d, "energy", "elec_mwh")
    parts = []
    for c, e in EXPECT_09.items():
        v = _q(ind[c] / mwh[c], 4)
        parts.append(f"{c} {v} (기대 {e}, 차이 {_q(v - Decimal(e), 4)})")
    return ["대조 — " + " · ".join(parts)]


def check_10(d):
    """검산 #10 — 원재료 원산지 증빙 없을 때 합 / 한국산 증빙 시 합 (EU 신고 표 precursor_t, 해당 경우 행)"""
    fails = []
    for case, e in EXPECT_10.items():
        got = _sum([_d(r["precursor_t"]) for r in d["eu_report"] if r["report_case"] == case])
        if not _same(Decimal(e), got):
            fails.append(_miss(e, got, case, "precursor_t"))
    return fails


def check_11(d):
    """검산 #11 — 생산 줄마다 12달 합 = annual_public"""
    fails = []
    for line, e in ANNUAL_PUBLIC.items():
        rows = [d["pr"][(line, m)] for m in MONTHS]
        got = _sum([_d(r["quantity"]) for r in rows])
        exp = None if e is None else Decimal(e)
        if not _same(exp, got):
            fails.append(_miss(exp, got, rows[0]["plant_id"], f"Σ quantity {line}번 줄"))
        for r in rows:
            if not _same(exp, _d(r["annual_public"])):
                fails.append(_miss(exp, r["annual_public"], r["plant_id"],
                                   f"annual_public {line}번 줄", r["month"]))
    return fails


def check_12(d):
    """검산 #12 — EU 신고 표 숫자 칸 (빈칸 자리도 같아야 함)"""
    head, rows = _md(TABLE_5)
    fails = []
    if len(d["eu_report"]) != len(rows):
        fails.append(_miss(len(rows), len(d["eu_report"]), "mfg_eu_report_synth.csv", "행 수"))
    for row in rows:
        r = d["eu"].get(row["row_no"])
        if r is None:
            fails.append(_miss(row["row_no"], None, "mfg_eu_report_synth.csv", "row_no"))
            continue
        if r["plant_id"] != row["plant_id"]:
            fails.append(_miss(row["plant_id"], r["plant_id"], f"row_no {row['row_no']}", "plant_id"))
        for h in head[3:]:
            exp, got = _exp(row[h]), _d(r[h])
            if not _same(exp, got):
                fails.append(_miss(exp, got, f"{row['plant_id']} row_no {row['row_no']}", h))
    return fails


def check_13(d):
    """검산 #13 — 은행 전달 표 숫자 칸, 회사 합계 줄 = 공장(묶음) 줄 합, total_t = 앵커"""
    head, rows = _md(TABLE_6)
    fails = []
    if len(d["bank"]) != len(rows):
        fails.append(_miss(len(rows), len(d["bank"]), "mfg_bank_transfer_synth.csv", "행 수"))
    by_plant = {}
    for row in rows:
        r = d["bk"].get(row["row_no"])
        if r is None:
            fails.append(_miss(row["row_no"], None, "mfg_bank_transfer_synth.csv", "row_no"))
            continue
        if r["plant_id"] != row["plant_id"]:
            fails.append(_miss(row["plant_id"], r["plant_id"], f"row_no {row['row_no']}", "plant_id"))
        by_plant[r["plant_id"]] = r
        for h in head[2:]:
            exp, got = _exp(row[h]), _d(r[h])
            if not _same(exp, got):
                fails.append(_miss(exp, got, row["plant_id"], h))
    # 회사 합계 줄 = 공장(묶음) 줄 합 — 양 칸 9개, 빈칸은 0으로 더함 (사용자 확정 2026-10-12)
    for total_id, parts in BANK_SUM_ROWS.items():
        if total_id not in by_plant or any(p not in by_plant for p in parts):
            fails.append(_miss(total_id, None, "mfg_bank_transfer_synth.csv", "plant_id"))
            continue
        for col in BANK_ADDITIVE:
            total = _d(by_plant[total_id][col]) or Decimal(0)
            s = sum((_d(by_plant[p][col]) or Decimal(0) for p in parts), Decimal(0))
            if total != s:
                fails.append(_miss(s, total, total_id, f"{col} (= {' + '.join(parts)})"))
    for total_id, e in BANK_ANCHOR.items():
        got = _d(by_plant.get(total_id, {}).get("total_t", ""))
        if not _same(Decimal(e), got):
            fails.append(_miss(e, got, total_id, "total_t = 앵커"))
    return fails


def check_14(d):
    """검산 #14 — 문구 대조 — plant_master 11행 eu_scope·eu_cn, eu_report 14행 report_case·value_label, bank 8행 value_label, energy 120행 envizi_input"""
    fails = []
    pm = {r["plant_id"]: r for r in d["plant_master"]}
    if len(pm) != len(PLANT_MASTER_TEXT):
        fails.append(_miss(len(PLANT_MASTER_TEXT), len(pm), "mfg_plant_master_synth.csv", "행 수"))
    for pid, (scope, cn) in PLANT_MASTER_TEXT.items():
        r = pm.get(pid, {})
        for col, exp in (("eu_scope", scope), ("eu_cn", cn)):
            if r.get(col) != exp:
                fails.append(_miss(exp or "빈칸", r.get(col), pid, col))
    for no, (case, label) in EU_REPORT_TEXT.items():
        r = d["eu"].get(no, {})
        for col, exp in (("report_case", case), ("value_label", label)):
            if r.get(col) != exp:
                fails.append(_miss(exp, r.get(col), f"eu_report row_no {no}", col))
    for no, label in BANK_TEXT.items():
        r = d["bk"].get(no, {})
        if r.get("value_label") != label:
            fails.append(_miss(label, r.get("value_label"), f"bank row_no {no}", "value_label"))
    for r in d["energy"]:
        if r["envizi_input"] != ENVIZI_TEXT:
            fails.append(_miss(ENVIZI_TEXT, r["envizi_input"], r["plant_id"], "envizi_input",
                               r["month"]))
    return fails


def check_15(d):
    """검산 #15 — data_nature 칸 빈 행 없음 (① 7개 파일, ③ CSV 3개), ① 파일 7개 이름이 모두 `_synth.csv`로 끝남"""
    fails = []
    for path, _, _, src in FILES:
        if not path.endswith(".csv"):
            continue
        name = Path(path).name
        if src == "①" and not name.endswith("_synth.csv"):
            fails.append(_miss("_synth.csv로 끝남", name, name, "파일 이름"))
        for i, r in enumerate(_read(ROOT / path), start=1):
            if not r.get("data_nature"):
                fails.append(_miss("값 있음", None, name, "data_nature", f"{i}행"))
    named = {Path(p).name for p, _, _, src in FILES if src == "①"}
    for f in sorted(DB.glob("*.csv")):
        if f.name not in named:
            fails.append(_miss("① 파일 7개만", f.name, "data/demo/mfg/db", "파일 이름"))
    return fails


def check_16(d):
    """검산 #16 — 행 수 (머리글 제외, 엑셀은 시트마다)"""
    fails = []
    for path, n, _, _ in FILES:
        name = Path(path).name
        if path.endswith(".xlsx"):
            wb = load_workbook(ROOT / path, data_only=True)
            if wb.sheetnames != list(SHEET_ROWS):
                fails.append(_miss(" · ".join(SHEET_ROWS), " · ".join(wb.sheetnames), name, "시트"))
            for sheet, e in SHEET_ROWS.items():
                if sheet not in wb.sheetnames:
                    continue
                got = sum(1 for row in wb[sheet].iter_rows(min_row=2, values_only=True)
                          if any(v is not None for v in row))
                if got != e:
                    fails.append(_miss(e, got, f"{name}#{sheet}", "행 수"))
        else:
            got = len(_read(ROOT / path))
            if got != n:
                fails.append(_miss(n, got, name, "행 수"))
    return fails


def check_17(d):
    """검산 #17 — 파일 크기"""
    lines = []
    for path, _, limit, _ in FILES:
        size = (ROOT / path).stat().st_size
        over = " 상한 넘음" if size > limit * KB else ""
        lines.append(f"파일 크기 — {path} {size} B (상한 {limit} KB){over}")
    return lines


CHECKS = [check_01, check_02, check_03, check_04, check_05, check_06, check_07, check_08,
          check_09, check_10, check_11, check_12, check_13, check_14, check_15, check_16,
          check_17]
INFO_CHECKS = {check_09, check_17}


def main():
    """검산 17개를 실행하고 요약 줄·종료 코드를 냄"""
    d = _load()
    counts = {"PASS": 0, "FAIL": 0, "INFO": 0}
    for i, check in enumerate(CHECKS, start=1):
        title = check.__doc__.splitlines()[0].split(" — ", 1)[1]
        out = check(d)
        if check in INFO_CHECKS:
            counts["INFO"] += 1
            for line in out:
                print(f"[INFO] #{i:02d} {line}")
        elif out:
            counts["FAIL"] += 1
            for line in out[:MAX_FAIL_LINES]:
                print(f"[FAIL] #{i:02d} {title} — {line}")
            if len(out) > MAX_FAIL_LINES:
                print(f"[FAIL] #{i:02d} {title} — 외 {len(out) - MAX_FAIL_LINES}칸")
        else:
            counts["PASS"] += 1
            print(f"[PASS] #{i:02d} {title}")
    print(f"PASS {counts['PASS']} · FAIL {counts['FAIL']} · INFO {counts['INFO']}")
    return 1 if counts["FAIL"] else 0


if __name__ == "__main__":
    sys.exit(main())
