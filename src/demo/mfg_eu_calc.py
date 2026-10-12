"""09 제조 시연 EU 신고 값 계산 (클로드 코드 1회차 — H4 ② 계산 구간)

사내 DB의 직접 배출·원재료 매입 수량·이동분과 계수 엑셀로 공장·품목별 EU 신고 값을 계산하고, 은행 전달 표를 만듦
규칙은 docs/progress/09_demo_mfg.md 4-8-8

입력
- data/demo/mfg/db/mfg_plant_master_synth.csv, mfg_energy_monthly_synth.csv, mfg_emissions_monthly_synth.csv, mfg_purchases_monthly_synth.csv, mfg_production_monthly_synth.csv
- data/demo/mfg/ext/mfg_ext_factors.xlsx (eu_defaults, gov_factors, markup)

출력
- data/demo/mfg/db/mfg_eu_report_synth.csv, mfg_bank_transfer_synth.csv
"""
# LINEAGE-STEP: mfg_eu_calc (H4 ② 계산 — 09 4-8-8)
# LINEAGE-INPUT: db/mfg_plant_master_synth.csv: company_id, plant_id, eu_scope, eu_cn
# LINEAGE-INPUT: db/mfg_energy_monthly_synth.csv: plant_id, fuel_tj_gross, elec_mwh
# LINEAGE-INPUT: db/mfg_emissions_monthly_synth.csv: plant_id, direct_t, indirect_t, total_t
# LINEAGE-INPUT: db/mfg_purchases_monthly_synth.csv: plant_id, default_cn_key, quantity_t
# LINEAGE-INPUT: db/mfg_production_monthly_synth.csv: line_no, quantity
# LINEAGE-INPUT: ext/mfg_ext_factors.xlsx#eu_defaults: cn, case, total_t_per_t
# LINEAGE-INPUT: ext/mfg_ext_factors.xlsx#gov_factors: factor_id, value
# LINEAGE-INPUT: ext/mfg_ext_factors.xlsx#markup: year, sector, rate
# LINEAGE-OUTPUT: db/mfg_eu_report_synth.csv: gov_total_t, own_t, precursor_t, embedded_t, production_t, see_t_per_t, precursor_x110_2026_t, embedded_markup_2026_t
# LINEAGE-OUTPUT: db/mfg_bank_transfer_synth.csv: direct_t, indirect_t, total_t, fuel_tj, elec_mwh, elec_factor_check, elec_factor_diff_pct, precursor_unknown_t, precursor_kr_t, eu_embedded_unknown_t, eu_embedded_kr_t
# LINEAGE-MAP: precursor_t = r3(sum(quantity_t * total_t_per_t))  (TRANSFER:TW_HJ 줄은 TW_HJ own_t / production_t)
# LINEAGE-MAP: embedded_t = own_t + precursor_t ; see_t_per_t = embedded_t / production_t
import sys
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mfg_common import DB, EXT, VERSION, r3, read_csv, rq, today_str, write_csv

COMMON_TAIL = ["data_nature", "method", "assumption", "source", "created_at", "version"]

# 아래 문구 표는 09 4-8-8 표를 줄 그대로 옮긴 것임. `빈칸` = 빈 값 (09 4-8-0)

EU_REPORT_TABLE = """
| row_no | plant_id | report_case | own_basis | precursor_detail | value_label | limits |
|---|---|---|---|---|---|---|
| 1 | KG_DJ | 원산지 증빙 없을 때 | 철강 — 직접 배출만 | 7208 1486328.623 t × 8.230 + 7210 176837.555 t × 8.250 | 원산지 불명 최고값 — 특정 공장 값 아님 | 석도 라인(plating 해석) 배출 포함 — 과대 (결정 24). KG 두 공장 1톤당 값 같음 (같은 비율로 나눔) |
| 2 | KG_DJ | 한국산 증빙 시 | 철강 — 직접 배출만 | 7208 1486328.623 t × 2.118 + 7210 176837.555 t × 2.144 | 국가 평균값 — 특정 공장 값 아님 | 석도 라인(plating 해석) 배출 포함 — 과대 (결정 24). KG 두 공장 1톤당 값 같음 (같은 비율로 나눔). 한국산 값의 가산 여부는 원문상 구분 없음 |
| 3 | KG_IC | 원산지 증빙 없을 때 | 철강 — 직접 배출만 | 7208 556216.377 t × 8.230 + 7210 66176.445 t × 8.250 | 원산지 불명 최고값 — 특정 공장 값 아님 | 석도 라인(plating 해석) 배출 포함 — 과대 (결정 24). KG 두 공장 1톤당 값 같음 (같은 비율로 나눔) |
| 4 | KG_IC | 한국산 증빙 시 | 철강 — 직접 배출만 | 7208 556216.377 t × 2.118 + 7210 66176.445 t × 2.144 | 국가 평균값 — 특정 공장 값 아님 | 석도 라인(plating 해석) 배출 포함 — 과대 (결정 24). KG 두 공장 1톤당 값 같음 (같은 비율로 나눔). 한국산 값의 가산 여부는 원문상 구분 없음 |
| 5 | TW_HJ | 해당 없음 (원재료 없음) | 철강 — 직접 배출만 (연소 + 전기로 공정) | 철 스크랩 7204 비대상, "외" 0 (결정 11) | 합성 실제값 (자체 공정) | 잉곳·블룸 대부분 녹산으로 이동 |
| 6 | TW_NS | 이동분 | 철강 — 직접 배출만 | 화전 잉곳·블룸 68165 t × 화전 1톤당 값 (25861.314 ÷ 180710) | 회사 안 이동 — 보내는 공장 합성 값 (EU 기본값 아님) | 일부 해당 — 비율 미공시, 녹산 전체 생산 기준 (결정 34). 외부 매입분 0 (결정 16) |
| 7 | HBG_CW | 원산지 증빙 없을 때 | 철강 — 직접 배출만 | 7219 239432 t × 8.670 | 원산지 불명 최고값 — 특정 공장 값 아님 | 원산지·강종 비율 미공시 |
| 8 | HBG_CW | 한국산 증빙 시 | 철강 — 직접 배출만 | 7219 239432 t × 3.630 | 국가 평균값 — 특정 공장 값 아님 | 원산지·강종 비율 미공시. 한국산 값의 가산 여부는 원문상 구분 없음 |
| 9 | HBG_SA | 이동 수량 미공시 | 자체 공정 배출 제외 (마무리 — 2025/2547 부속서 I 3.16.2) | 창원 판 이동분 — 수량 미공시 | 이동 수량 미공시 | EU 대상 물건(추정) (결정 23) |
| 10 | HBG_GJ | 비해당 | 빈칸 | 빈칸 | 비해당 — 정부 신고에만 있음 | 자동차 엔진 블럭용 사형코어 |
| 11 | SA_PS | 원산지 증빙 없을 때 | 알루미늄 — 직접 배출만 | 7606 15457 t × 4.986 | 원산지 불명 최고값 — 특정 공장 값 아님 | 가공(finishing 해석) 배출 포함 — 과대 (결정 24). 수량 상한 17,910 t이면 원재료 최대 약 16% 큼 |
| 12 | SA_PS | 한국산 증빙 시 | 알루미늄 — 직접 배출만 | 7606 15457 t × 1.950 | 국가 평균값 — 특정 공장 값 아님 | 가공(finishing 해석) 배출 포함 — 과대 (결정 24). 수량 상한 17,910 t이면 원재료 최대 약 16% 큼. 한국산 값의 가산 여부는 원문상 구분 없음 |
| 13 | HI_CEM | 해당 없음 (원재료 없음) | 시멘트 — 직접 + 간접 | 클링커 자체 생산 — 매입 0 (결정 11) | 회사 사업별 공개값 (자체 공정) | 당진·평택 고로슬래그 미분말(비대상) 몫 섞임, 생산량에 슬래그파우더 포함 → 1톤당 값 낮음 (결정 22) |
| 14 | HI_SEC | 비해당 | 빈칸 | 빈칸 | 비해당 — 정부 신고에만 있음 | 레미콘·레미탈 |
"""
EU_REPORT_COLUMNS = ["row_no", "company_id", "plant_id", "eu_scope", "eu_cn", "report_case",
                     "own_basis", "gov_total_t", "own_t", "precursor_t", "precursor_detail",
                     "embedded_t", "production_t", "see_t_per_t", "precursor_x110_2026_t",
                     "embedded_markup_2026_t", "value_label", "limits",
                     "calc_step"] + COMMON_TAIL
EU_REPORT_ASSUMPTION = {1: "가정: 원재료 공장 비율 (결정 33)", 2: "가정: 원재료 공장 비율 (결정 33)",
                        3: "가정: 원재료 공장 비율 (결정 33)", 4: "가정: 원재료 공장 비율 (결정 33)",
                        6: "가정: 받은 양 = 만든 양, 이동분 값 = 화전 1톤당 값 (결정 46)",
                        7: "가정: 사 온 양 = 만든 양", 8: "가정: 사 온 양 = 만든 양",
                        11: "가정: 사 온 양 = 만든 양", 12: "가정: 사 온 양 = 만든 양"}
EU_REPORT_NATURE = "합성 (Python 단계 결과)"
EU_REPORT_METHOD = "09 4-8-8 계산식"
EU_REPORT_SOURCE = "09 4-7-6·4-8-8"
CALC_STEP = "src/demo/mfg_eu_calc.py"

# own_t 규칙 (09 4-8-8): 철강·알루미늄은 직접 배출만, 시멘트는 직접 + 간접, 송악은 마무리 공정이라 0, 비해당은 빈칸
OWN_DIRECT = {"KG_DJ", "KG_IC", "TW_HJ", "TW_NS", "HBG_CW", "SA_PS"}
OWN_TOTAL = {"HI_CEM"}
OWN_ZERO = {"HBG_SA"}
# production_t = 생산 파일 이 줄의 Σ quantity (CBAM 기능 단위 = t)
PRODUCTION_LINE = {"KG_DJ": "1", "KG_IC": "2", "TW_HJ": "3", "TW_NS": "4", "HBG_CW": "5",
                   "SA_PS": "8", "HI_CEM": "10"}
# report_case → eu_defaults case
DEFAULT_CASE = {"원산지 증빙 없을 때": "원산지 불명", "한국산 증빙 시": "한국"}
CASE_TRANSFER = "이동분"
CASE_NONE = "해당 없음 (원재료 없음)"
TRANSFER_PREFIX = "TRANSFER:"
MARKUP_SECTOR = "시멘트·철강·알루미늄·수소"
MARKUP_YEAR = 2026

BANK_TABLE = """
| row_no | company_id | linked_bank | row_kind | plant_id | production_text | verifier | value_label | data_nature |
|---|---|---|---|---|---|---|---|---|
| 1 | KG | 하나은행 | 회사 합계 | KG_ALL | 판재 2207000 t | (재)한국품질재단 | 합계 = 실측(명세서 2025), 직접·간접 = 합성 — 실제로는 회사 실측 자리 | 합성 (Python 단계 결과) |
| 2 | KG | 하나은행 | 공장 | KG_DJ | 판재 1606000 t | (재)한국품질재단 (회사 합계 검증) | 합성 — 실제로는 회사 실측 자리 | 합성 (Python 단계 결과) |
| 3 | KG | 하나은행 | 공장 | KG_IC | 판재 601000 t | (재)한국품질재단 (회사 합계 검증) | 합성 — 실제로는 회사 실측 자리 | 합성 (Python 단계 결과) |
| 4 | HI | 한국산업은행 | 회사 합계 | HI_ALL | 시멘트 외 9260000 t · 레미콘 2022000 ㎥ · 레미탈 3222000 t | 한국표준협회 | 실측 (명세서 2025 = 지속가능경영보고서) | 합성 (Python 단계 결과) |
| 5 | HI | 한국산업은행 | 사업 묶음 | HI_CEM | 시멘트 외 9260000 t | 한국표준협회 (회사 합계 검증) | 회사 공개값 (배출) · 합성 (에너지 — 결정 39) | 합성 (Python 단계 결과) |
| 6 | HI | 한국산업은행 | 사업 묶음 | HI_SEC | 레미콘 2022000 ㎥ · 레미탈 3222000 t | 한국표준협회 (회사 합계 검증) | 회사 공개값 (배출) · 합성 (에너지 — 결정 39) | 합성 (Python 단계 결과) |
| 7 | SA | 한국수출입은행 | 회사 합계 | SA_ALL | 압연 15457 t · 가공 118955 km | 한국표준협회 | 합계 = 실측(명세서 2025), 직접·간접 = 합성 — 실제로는 회사 실측 자리 | 합성 (Python 단계 결과) |
| 8 | SA | 한국수출입은행 | 공장 | SA_PS | 압연 15457 t · 가공 118955 km | 한국표준협회 (회사 합계 검증) | 합성 — 실제로는 회사 실측 자리 | 합성 (Python 단계 결과) |
"""
BANK_COLUMNS = ["row_no", "company_id", "company_name", "linked_bank", "row_kind", "plant_id",
                "year", "direct_t", "indirect_t", "total_t", "fuel_tj", "elec_mwh",
                "production_text", "verifier", "method_basis", "elec_factor_check",
                "elec_factor_diff_pct", "precursor_unknown_t", "precursor_kr_t",
                "eu_embedded_unknown_t", "eu_embedded_kr_t", "value_label"] + COMMON_TAIL
BANK_COMPANY_NAME = {"KG": "KG스틸", "HI": "한일시멘트", "SA": "삼아알미늄"}
BANK_YEAR = 2025
BANK_METHOD_BASIS = {
    "KG": "온실가스 배출권거래제 보고·인증 지침 별표6 — 전기 0.4594106 tCO2eq/MWh, 연료 56.1 tCO2/TJ(순발열량, 도시가스 가정), 명세서 에너지 = 총발열량 가정",
    "HI": "온실가스 배출권거래제 보고·인증 지침 — 회사 공개값(지속가능경영보고서 ESG Databook, 배출권 제출 기준). 묶음 에너지는 결정 39",
    "SA": "온실가스 목표관리 운영 등에 관한 지침 (목표관리업체 — 배출권 지침과 같은 산정 방법으로 봄, 가정) — 전기 0.4594106 tCO2eq/MWh, 연료 56.1 tCO2/TJ(순발열량, 도시가스 가정)",
}
BANK_METHOD = "db 월별 값의 연간 합 (09 4-8-8)"
BANK_SOURCE = "09 4-7-7·4-8-8"
# assumption: KG 줄·SA 줄·HI 묶음 줄 (HI 회사 합계 줄은 빈칸)
BANK_ASSUMPTION = {"KG": "가정: 직접·간접 = 연료 먼저 (결정 28), 공장 = 생산실적 비율 (결정 1)",
                   "SA": "가정: 직접·간접 = 연료 먼저 (결정 28)",
                   "HI_CEM": "가정: 묶음 에너지 (결정 39)", "HI_SEC": "가정: 묶음 에너지 (결정 39)"}
# 은행 전달 표 줄 → 더하는 공장 (회사 합계 줄 = 그 회사 공장 전부)
BANK_PLANTS = {"KG_ALL": ["KG_DJ", "KG_IC"], "HI_ALL": ["HI_CEM", "HI_SEC"],
               "SA_ALL": ["SA_PS"]}
# 한일은 원재료 매입 없음 (결정 11) — 원재료 칸 0, EU 값은 EU 신고 표 HI_CEM 행 (HI_SEC 비해당)
HI_ZERO_PRECURSOR = "HI"
HI_EU_ROW_PLANT = "HI_CEM"
HI_EU_BANK_ROWS = {"HI_ALL", "HI_CEM"}
DECIMALS = {"see_t_per_t": 4, "elec_factor_check": 4, "elec_factor_diff_pct": 1}


def _md_rows(text):
    """마크다운 표(머리글 줄 + 구분 줄 + 값 줄)를 칸 이름 → 값 사전 목록으로 바꿈 (09 4-8-0 표 읽는 법)"""
    lines = [ln.strip() for ln in text.strip().splitlines()]
    head = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = []
    for ln in lines[2:]:
        cells = [c.strip() for c in ln.strip("|").split("|")]
        if len(cells) != len(head):
            raise ValueError(f"표 칸 수가 머리글과 다름: {ln}")
        if rows:
            cells = [rows[-1][i] if c == "같음" else c for i, c in enumerate(cells)]
        rows.append(cells)
    return [{h: (None if c == "빈칸" else c) for h, c in zip(head, r)} for r in rows]


def _add(total, text):
    """12달 합 — 빈 문자열은 값 없음. 모두 비면 None, 일부만 비면 멈춤"""
    if text == "":
        if total not in (None, "blank"):
            raise ValueError("12달 중 일부만 빈칸 — 멈춤")
        return "blank"
    if total == "blank":
        raise ValueError("12달 중 일부만 빈칸 — 멈춤")
    return Fraction(text) + (total or 0)


def load_db():
    """입력: db 5개 파일 / 출력: 공장·줄별 12달 합"""
    master = {r["plant_id"]: r for r in read_csv(DB / "mfg_plant_master_synth.csv")}
    plants = {}
    for name, cols in (("mfg_emissions_monthly_synth.csv", ["direct_t", "indirect_t", "total_t"]),
                       ("mfg_energy_monthly_synth.csv", ["fuel_tj_gross", "elec_mwh"])):
        for r in read_csv(DB / name):
            p = plants.setdefault(r["plant_id"], {})
            for c in cols:
                p[c] = _add(p.get(c), r[c])
    purchases = {}
    for r in read_csv(DB / "mfg_purchases_monthly_synth.csv"):
        p = purchases.setdefault(r["line_no"], {"plant_id": r["plant_id"],
                                                "default_cn_key": r["default_cn_key"],
                                                "quantity_t": None})
        p["quantity_t"] = _add(p["quantity_t"], r["quantity_t"])
    production = {}
    for r in read_csv(DB / "mfg_production_monthly_synth.csv"):
        production[r["line_no"]] = _add(production.get(r["line_no"]), r["quantity"])
    for d in [*plants.values(), *purchases.values()]:
        for k, v in d.items():
            if v == "blank":
                d[k] = None
    production = {k: (None if v == "blank" else v) for k, v in production.items()}
    return {"master": master, "plants": plants, "purchases": purchases,
            "production": production}


def load_factors():
    """입력: 엑셀 eu_defaults·gov_factors·markup / 출력: 기본값·계수·가산 비율"""
    wb = load_workbook(EXT / "mfg_ext_factors.xlsx", data_only=True)

    def sheet(name):
        it = wb[name].iter_rows(values_only=True)
        head = [str(c) for c in next(it)]
        return [dict(zip(head, row)) for row in it]

    defaults = {(str(r["cn"]), str(r["case"])): Fraction(str(r["total_t_per_t"]))
                for r in sheet("eu_defaults")}
    gov = {str(r["factor_id"]): Fraction(str(r["value"])) for r in sheet("gov_factors")}
    rates = [Fraction(str(r["rate"])) for r in sheet("markup")
             if r["year"] == MARKUP_YEAR and str(r["sector"]) == MARKUP_SECTOR]
    if len(rates) != 1:
        raise ValueError("markup 시트에서 2026 시멘트·철강·알루미늄·수소 행이 1개가 아님 — 멈춤")
    return {"defaults": defaults, "gov": gov, "rate": rates[0]}


def calc_eu_report(db, factors):
    """입력: 공장별 합·기본값·가산 비율 / 출력: eu_report 14행"""
    plants, purchases, production = db["plants"], db["purchases"], db["production"]
    defaults, rate = factors["defaults"], factors["rate"]

    def own(pid):
        if pid in OWN_DIRECT:
            return plants[pid]["direct_t"]
        if pid in OWN_TOTAL:
            return plants[pid]["total_t"]
        if pid in OWN_ZERO:
            return Fraction(0)
        return None

    def prod(pid):
        line = PRODUCTION_LINE.get(pid)
        return None if line is None else production[line]

    rows = []
    for r in _md_rows(EU_REPORT_TABLE):
        pid, case = r["plant_id"], r["report_case"]
        lines = [p for p in purchases.values() if p["plant_id"] == pid]
        own_t = own(pid)
        if case in DEFAULT_CASE:
            keyed = [p for p in lines if p["default_cn_key"]
                     and not p["default_cn_key"].startswith(TRANSFER_PREFIX)]
            precursor = r3(sum(p["quantity_t"] * defaults[(p["default_cn_key"], DEFAULT_CASE[case])]
                               for p in keyed))
        elif case == CASE_TRANSFER:
            precursor = Fraction(0)
            for p in lines:
                if p["default_cn_key"].startswith(TRANSFER_PREFIX):
                    sender = p["default_cn_key"][len(TRANSFER_PREFIX):]
                    # 보내는 공장 1톤당 값 = 그 공장 own_t ÷ production_t (반올림하지 않은 분수)
                    precursor += p["quantity_t"] * (own(sender) / prod(sender))
            precursor = r3(precursor)
        elif case == CASE_NONE:
            precursor = Fraction(0)
        else:
            precursor = None
        embedded = own_t + precursor if own_t is not None and precursor is not None else None
        production_t = prod(pid)
        see = (rq(embedded / production_t, 4)
               if embedded is not None and production_t is not None else None)
        x110 = r3(precursor * (1 + rate)) if case in DEFAULT_CASE else None
        markup = own_t + x110 if x110 is not None else None
        m = db["master"][pid]
        row_no = int(r["row_no"])
        rows.append({"row_no": row_no, "company_id": m["company_id"], "plant_id": pid,
                     "eu_scope": m["eu_scope"] or None, "eu_cn": m["eu_cn"] or None,
                     "report_case": case, "own_basis": r["own_basis"],
                     "gov_total_t": plants[pid]["total_t"], "own_t": own_t,
                     "precursor_t": precursor, "precursor_detail": r["precursor_detail"],
                     "embedded_t": embedded, "production_t": production_t,
                     "see_t_per_t": see, "precursor_x110_2026_t": x110,
                     "embedded_markup_2026_t": markup, "value_label": r["value_label"],
                     "limits": r["limits"], "calc_step": CALC_STEP,
                     "data_nature": EU_REPORT_NATURE, "method": EU_REPORT_METHOD,
                     "assumption": EU_REPORT_ASSUMPTION.get(row_no), "source": EU_REPORT_SOURCE})
    return rows


def calc_bank_transfer(db, report, factors):
    """입력: 공장별 합·eu_report·ELEC_EF / 출력: bank_transfer 8행"""
    plants = db["plants"]
    elec_ef = factors["gov"]["ELEC_EF"]
    rows = []
    for r in _md_rows(BANK_TABLE):
        cid, bid = r["company_id"], r["plant_id"]
        targets = BANK_PLANTS.get(bid, [bid])
        s = {c: sum(plants[p][c] for p in targets)
             for c in ("direct_t", "indirect_t", "total_t", "fuel_tj_gross", "elec_mwh")}
        ratio = s["indirect_t"] / s["elec_mwh"]
        if cid == HI_ZERO_PRECURSOR:
            p_unknown = p_kr = Fraction(0)
            hi = next(x for x in report if x["plant_id"] == HI_EU_ROW_PLANT)
            e_unknown = e_kr = hi["embedded_t"] if bid in HI_EU_BANK_ROWS else None
        else:
            unknown = [x for x in report if x["plant_id"] in targets
                       and x["report_case"] == "원산지 증빙 없을 때"]
            kr = [x for x in report if x["plant_id"] in targets
                  and x["report_case"] == "한국산 증빙 시"]
            p_unknown = sum(x["precursor_t"] for x in unknown)
            p_kr = sum(x["precursor_t"] for x in kr)
            e_unknown = sum(x["embedded_t"] for x in unknown)
            e_kr = sum(x["embedded_t"] for x in kr)
        rows.append({"row_no": int(r["row_no"]), "company_id": cid,
                     "company_name": BANK_COMPANY_NAME[cid], "linked_bank": r["linked_bank"],
                     "row_kind": r["row_kind"], "plant_id": bid, "year": BANK_YEAR,
                     "direct_t": s["direct_t"], "indirect_t": s["indirect_t"],
                     "total_t": s["total_t"], "fuel_tj": s["fuel_tj_gross"],
                     "elec_mwh": s["elec_mwh"], "production_text": r["production_text"],
                     "verifier": r["verifier"], "method_basis": BANK_METHOD_BASIS[cid],
                     "elec_factor_check": rq(ratio, 4),
                     "elec_factor_diff_pct": rq((ratio / elec_ef - 1) * 100, 1),
                     "precursor_unknown_t": p_unknown, "precursor_kr_t": p_kr,
                     "eu_embedded_unknown_t": e_unknown, "eu_embedded_kr_t": e_kr,
                     "value_label": r["value_label"], "data_nature": r["data_nature"],
                     "method": BANK_METHOD,
                     "assumption": BANK_ASSUMPTION.get(cid, BANK_ASSUMPTION.get(bid)),
                     "source": BANK_SOURCE})
    return rows


def write_results(report, bank, created):
    """출력: eu_report·bank_transfer CSV"""
    for name, columns, rows in (("mfg_eu_report_synth.csv", EU_REPORT_COLUMNS, report),
                                ("mfg_bank_transfer_synth.csv", BANK_COLUMNS, bank)):
        rows = [{**r, "created_at": created, "version": VERSION} for r in rows]
        write_csv(DB / name, columns, rows, DECIMALS)
        print(f"{name}: {len(rows)}행")


def main():
    """EU 신고 값 계산을 차례로 실행함"""
    created = today_str()
    db = load_db()
    factors = load_factors()
    report = calc_eu_report(db, factors)
    bank = calc_bank_transfer(db, report, factors)
    write_results(report, bank, created)


if __name__ == "__main__":
    main()
