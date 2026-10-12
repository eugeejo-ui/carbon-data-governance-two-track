"""09 제조 시연 ① 제조사 사내 DB 자리 합성 (클로드 코드 1회차)

앵커(명세서 2025)를 공개 통계 모양으로 공장·월에 나누고, 원재료 매입 수량을 합성함 (H4 시험 대상 아님)
규칙은 docs/progress/09_demo_mfg.md 4-8-6·4-8-7

입력
- data/demo/mfg/ext/mfg_ext_anchor.csv, mfg_ext_factors.xlsx(gov_factors), mfg_ext_stats.csv, mfg_ext_disclosed.csv

출력
- data/demo/mfg/db/mfg_plant_master_synth.csv, mfg_energy_monthly_synth.csv, mfg_production_monthly_synth.csv, mfg_emissions_monthly_synth.csv, mfg_purchases_monthly_synth.csv
"""
# LINEAGE-STEP: mfg_synth_build (H4 시험 대상 아님 — 합성 생성)
# LINEAGE-INPUT: ext/mfg_ext_anchor.csv: company_id, emissions_tco2eq, energy_tj
# LINEAGE-INPUT: ext/mfg_ext_factors.xlsx#gov_factors: factor_id, value
# LINEAGE-INPUT: ext/mfg_ext_stats.csv: stat_id, month, value
# LINEAGE-INPUT: ext/mfg_ext_disclosed.csv: company_id, item_id, value
# LINEAGE-OUTPUT: db/mfg_plant_master_synth.csv: 전 칸
# LINEAGE-OUTPUT: db/mfg_energy_monthly_synth.csv: 전 칸
# LINEAGE-OUTPUT: db/mfg_production_monthly_synth.csv: 전 칸
# LINEAGE-OUTPUT: db/mfg_emissions_monthly_synth.csv: 전 칸
# LINEAGE-OUTPUT: db/mfg_purchases_monthly_synth.csv: 전 칸
# LINEAGE-MAP: 앵커 → 공장 연간(4-8-7 ①) → 월(4-8-7 ②, 최대 나머지 배분)
import sys
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mfg_common import (DB, EXT, MONTHS, PLANTS, VERSION, lr, r0, r3, read_csv,
                        today_str, write_csv)

COMMON_TAIL = ["data_nature", "method", "assumption", "source", "created_at", "version"]

# 아래 문구 표는 09 4-8-6 표를 줄 그대로 옮긴 것임. `빈칸` = 빈 값, `같음` = 바로 위 줄 같은 칸 (09 4-8-0)
# 숫자 칸(생산 annual_public·매입 연간 quantity_t)은 옮기지 않음 — ③ 파일에서 읽어 계산함 (09 4-8-7)

PLANT_MASTER_TABLE = """
| company_id | company_name | plant_id | plant_name | plant_kind | sido | elec_shape_id | prod_shape_id | eu_process | eu_scope | eu_cn | eu_note | register_basis | assumption |
|---|---|---|---|---|---|---|---|---|---|---|---|---|---|
| KG | KG스틸 | KG_DJ | 당진 | 공장 | 충남 | M1_CHUNGNAM | P_COLD | KG 당진 — 철강 제품 | 해당 — 석도 라인(plating 해석) 배출 포함 | 7209·7210 | 제품별 구분 없음 (결정 26) | 등록공장 "냉연(아연도, 석도) 강판", 사업보고서 사업소 당진공장 | 빈칸 |
| KG | KG스틸 | KG_IC | 인천 | 공장 | 인천 | M1_INCHEON | P_COLD | KG 인천 — 철강 제품 | 해당 — 석도 라인(plating 해석) 배출 포함 | 7209·7210 | 제품별 구분 없음 (결정 26) | 등록공장 "냉간압연압출및인발제품제조 외 6", 사업보고서 사업소 인천공장 | 빈칸 |
| TW | 태웅 | TW_HJ | 화전 | 공장 | 부산 | M1_BUSAN | P_STEEL1 | 태웅 화전 — 조강 | 해당 — 잉곳·블룸 (대부분 녹산으로 이동) | 7206·7207·7218 등 (품목 추정) | 제강사업부 | 등록공장 "인고트, BLOOM", 사업보고서 지점 | 빈칸 |
| TW | 태웅 | TW_NS | 녹산 | 공장 | 부산 | M1_BUSAN | P_STEEL1 | 태웅 녹산 — 철강 제품 | 일부 해당 — 비율 미공시 | 7308(타워 플랜지 — 대법원 2011두12290) / 7326 또는 84류(그 외 — 불명) | EU 대상 확인 제품(타워 플랜지)이 속한 풍력 부품 매출 49.1%(2025 사업보고서) — 플랜지 몫은 이보다 작음 | 등록공장 "단조품 / 플랜지 등 단조품 가공", 사업보고서 본점(단조사업부) | 빈칸 |
| TW | 태웅 | TW_DD | 다대 | 공장(숫자 줄 없음) | 부산 | 빈칸 | 빈칸 | 빈칸 | 숫자 줄 없음 (사업부 단위로 나눔) | 불명 | 다대 몫은 화전·녹산 숫자 안에 있는 것으로 봄 (결정 42) | 사업보고서 사업장 목록에만 있음(사하구 다산로176번길 30), 등록공장 명단에 없음 | 가정: 다대 몫은 화전·녹산 숫자 안에 있음 |
| HBG | 현대비앤지스틸 | HBG_CW | 창원 | 공장 | 경남 | M1_GYEONGNAM | P_COLD | 현대비앤지 창원 — 철강 제품 | 해당 | 7219·7220 | 폭 구분 미공시 (기본값 같음) | 등록공장 "스텐레스·강판·강대", 사업보고서 생산능력·실적 기준 공장 | 빈칸 |
| HBG | 현대비앤지스틸 | HBG_GJ | 광주 | 공장 | 광주 | M1_GWANGJU | P_COLD | 빈칸 | 비해당 — 자동차 엔진 블럭용 사형코어 | 빈칸 | 정부 신고에만 있음 | 등록공장 "자동차엔진부품", 사업보고서 생산시설 | 빈칸 |
| HBG | 현대비앤지스틸 | HBG_SA | 송악 | 공장 | 충남 | M1_CHUNGNAM | P_COLD | 빈칸 | EU 대상 물건(추정) — 자체 공정 배출 제외(마무리) | 7219·7220 (추정) | EU 값 = 창원 판 이동분, 이동 수량 미공시 (결정 23·33) | 등록공장 "철강 가공품", 2024 지속가능경영보고서 "스테인리스 표면 가공품" 3.1만 t/년 | 가정: 송악 제품은 스테인리스 판 (생산품 미공시) |
| SA | 삼아알미늄 | SA_PS | 포승 | 공장 | 경기 | M1_GYEONGGI | P_ALFOIL | 삼아 포승 — 알루미늄 제품 | 해당 — 가공(라미네이팅·인쇄, finishing 해석) 배출 포함 | 7607 | 빈칸 | 등록공장 "압연·알페이스트·가공지·인쇄", 사업보고서 포승공장 | 빈칸 |
| HI | 한일시멘트 | HI_CEM | 시멘트 묶음 | 사업 묶음 | 충북·강원·충남 | M1_CEMENT_MIX | P_CEMENT | 한일 시멘트 묶음 — 시멘트 | 해당 — 단양·삼곡·영월 포틀랜드, 당진·평택 일부 해당(슬래그시멘트 해당, 고로슬래그 미분말 비해당) — 비대상 미분말 몫 섞임 | 2523 | 공장 5곳 묶음 — 공장별 값 없음 (결정 20·22) | 단양·삼곡(충북), 영월(강원), 당진·평택(충남 — 평택 등록 주소 당진) | 가정: 월별 전력 = 충북·강원·충남 가중 모양 (결정 21) |
| HI | 한일시멘트 | HI_SEC | 2차제품 묶음 | 사업 묶음 | 전국 | M1_NATION | P_CEMENT | 빈칸 | 비해당 — 레미콘·레미탈 (2026/1740 한국 표에 없음) | 빈칸 | 정부 신고에만 있음 | 레미콘 10곳·레미탈 8곳 (사업보고서 p.32) | 대신 쓴 자료: 전국 전력 모양·시멘트 산업생산지수 (결정 14) |
"""
PLANT_MASTER_COLUMNS = ["company_id", "company_name", "plant_id", "plant_name", "plant_kind",
                        "sido", "elec_shape_id", "prod_shape_id", "eu_process", "eu_scope",
                        "eu_cn", "eu_note", "register_basis"] + COMMON_TAIL
PLANT_MASTER_NATURE = "합성"
PLANT_MASTER_METHOD = "공개 자료(등록공장·사업보고서·지속가능경영보고서)로 구성한 사내 공장 목록 대역"
PLANT_MASTER_SOURCE = "09 4-5-6·4-7-5"

ENERGY_TABLE = """
| plant_id | ratio_basis | assumption |
|---|---|---|
| KG_DJ | 광업제조업조사 2024 24122 사용량 환산 — 전기 74.6752% (결정 5·7) | 가정: 연료 = 도시가스, 명세서 에너지 = 총발열량, 연료 월 = 생산 모양 (결정 17·29). 대신 쓴 자료: 냉연 월 모양 |
| KG_IC | 같음 | 같음 |
| TW_HJ | 광업제조업조사 2024 24112 사용량 환산 — 전기 80.7447% (결정 5·7) | 가정: 연료 = 도시가스, 명세서 에너지 = 총발열량, 연료 월 = 생산 모양 (결정 17·29). 대신 쓴 자료: 1차 철강 생산지수 |
| TW_NS | 광업제조업조사 2024 25912 사용량 환산 — 전기 66.4738% (결정 5·7) | 같음 |
| HBG_CW | 회사 공개값 — 전기 1,094 / (연료 235 + 전기 1,094) TJ = 82.3175% (결정 1·27) | 가정: 연료 월 = 생산 모양 (결정 17). 대신 쓴 자료: 냉연 월 모양 |
| HBG_GJ | 같음 | 가정: 연료 월 = 생산 모양 (결정 17). 대신 쓴 자료: 냉연 월 모양 — 회사 주력 제품 모양 (결정 40) |
| HBG_SA | 같음 | 같음 |
| SA_PS | 광업제조업조사 2024 24222 사용량 환산 — 전기 65.2965% (결정 5·7) | 가정: 연료 = 도시가스, 명세서 에너지 = 총발열량, 연료 월 = 생산 모양 (결정 17·29) |
| HI_CEM | 전기 = 일반 전력 9,319 TJ를 사업별 Scope2(370,165 : 40,418)로 나눔, 연료 = 32,657 TJ − 2차제품 연료 (결정 39) | 가정: 연료 월 = 생산 모양 (결정 17), 전력 월 = 충북·강원·충남 가중 모양 (결정 21) |
| HI_SEC | 전기 = 일반 전력 9,319 TJ를 사업별 Scope2(370,165 : 40,418)로 나눔, 연료 = 2차제품 Scope1 54,947 ÷ 56.1 × 42.7 ÷ 38.5 (결정 39) | 가정: 2차제품 연료 = 도시가스, 연료 월 = 생산 모양 (결정 17·39). 대신 쓴 자료: 전국 전력 모양·시멘트 산업생산지수 (결정 14) |
"""
ENERGY_COLUMNS = ["month", "company_id", "plant_id", "fuel_tj_gross", "fuel_tj_net", "elec_tj",
                  "elec_mwh", "energy_tj", "ratio_basis", "fuel_shape_id", "elec_shape_id",
                  "envizi_input"] + COMMON_TAIL
ENERGY_ENVIZI = "Envizi 입력값: elec_mwh, fuel_tj_gross"
ENERGY_NATURE = "합성"
ENERGY_METHOD = "연간 공장 값(4-8-7 ①) → 전기는 M1 모양, 연료는 제품 모양으로 최대 나머지 배분 (0.001 TJ)"
ENERGY_SOURCE = "09 4-7-2·4-7-3·4-8-7"

PRODUCTION_TABLE = """
| line_no | plant_id | product | eu_cn | unit | prod_shape_id | assumption | source |
|---|---|---|---|---|---|---|---|
| 1 | KG_DJ | 판재 (냉연·아연도·칼라·석도) | 7209·7210 | t | P_COLD | 대신 쓴 자료: 냉연 월 모양 (결정 26 — 모든 판이 냉연 라인을 거침) | KG스틸 2025 사업보고서 II-3 생산실적 |
| 2 | KG_IC | 판재 (냉연·아연도·칼라·석도) | 7209·7210 | t | P_COLD | 같음 | 같음 |
| 3 | TW_HJ | 잉곳·블룸 (제강) | 7206·7207·7218 등 (품목 추정) | t | P_STEEL1 | 대신 쓴 자료: 1차 철강 생산지수 (결정 8) | 태웅 2025 사업보고서 II-3 생산실적 |
| 4 | TW_NS | 단조품 | 7308(타워 플랜지 — 대법원 2011두12290) / 7326 또는 84류(그 외 — 불명) | t | P_STEEL1 | 같음 | 같음 |
| 5 | HBG_CW | 스테인리스 강판 | 7219·7220 | t | P_COLD | 대신 쓴 자료: 냉연(스테인리스 포함) 월 모양 (결정 8) | 현대비앤지스틸 2025 사업보고서 II-3 (p.16~17) |
| 6 | HBG_GJ | 자동차부품 등 | 빈칸 | 개 | P_COLD | 대신 쓴 자료: 냉연 월 모양 — 회사 주력 제품 모양 (결정 40) | 같음 |
| 7 | HBG_SA | 스테인리스 표면 가공품 | 7219·7220 (추정) | t | 빈칸 | 빈칸 | 현대비앤지스틸 2024 지속가능경영보고서 (능력 3.1만 t/년만) |
| 8 | SA_PS | 압연 | 7607 | t | P_ALFOIL | 빈칸 | 삼아알미늄 2025 사업보고서 (p.12~14) |
| 9 | SA_PS | 가공 | 7607 | km | P_ALFOIL | 대신 쓴 자료: 알루미늄박 월 모양 (결정 14) | 같음 |
| 10 | HI_CEM | 시멘트 외 (시멘트 + 슬래그파우더) | 2523 | t | P_CEMENT | 빈칸 | 한일시멘트 2025 사업보고서 p.32 |
| 11 | HI_SEC | 레미콘 | 빈칸 | ㎥ | P_CEMENT | 대신 쓴 자료: 시멘트 산업생산지수 (결정 14) | 같음 |
| 12 | HI_SEC | 레미탈 | 빈칸 | t | P_CEMENT | 같음 | 같음 |
"""
PRODUCTION_COLUMNS = ["month", "company_id", "plant_id", "line_no", "product", "eu_cn",
                      "quantity", "unit", "annual_public", "prod_shape_id"] + COMMON_TAIL
# 연간 공개 생산실적을 읽는 곳 (09 4-8-7 숫자를 읽는 곳): 줄 → (회사, 항목, 단위 바꾸기 — 천t·천㎥ × 1,000)
PRODUCTION_ITEM = {"1": ("KG", "PROD_DJ", 1000), "2": ("KG", "PROD_IC", 1000),
                   "3": ("TW", "PROD_STEEL", 1), "4": ("TW", "PROD_FORGE", 1),
                   "5": ("HBG", "PROD_CW", 1), "6": ("HBG", "PROD_GJ", 1), "7": None,
                   "8": ("SA", "PROD_ROLL", 1), "9": ("SA", "PROD_PROC", 1),
                   "10": ("HI", "PROD_CEM", 1000), "11": ("HI", "PROD_REMICON", 1000),
                   "12": ("HI", "PROD_REMITAL", 1000)}
PRODUCTION_NATURE = "합성"
PRODUCTION_METHOD = "연간 공개 생산실적을 제품 모양으로 최대 나머지 배분"
PRODUCTION_METHOD_BLANK = "공란 — 생산량 미공시 (능력 3.1만 t/년만 공개)"

EMISSIONS_TABLE = """
| plant_id | split_basis | assumption |
|---|---|---|
| KG_DJ | 연료 먼저 — 직접 = 연료 TJ × 38.5 ÷ 42.7 × 56.1, 간접 = 앵커 − 직접 (결정 28·29) | 가정: 연료 = 도시가스, 명세서 에너지 = 총발열량, 직접 배출 월 = 생산 모양 (결정 17) |
| KG_IC | 같음 | 같음 |
| TW_HJ | 연료 먼저 + 전기로 공정 180,710 t × 0.08 (결정 28) | 같음 |
| TW_NS | 연료 먼저 — 직접 = 연료 TJ × 38.5 ÷ 42.7 × 56.1, 간접 = 앵커 − 직접 (결정 28·29) | 같음 |
| HBG_CW | 회사 공개 비율 Scope1 12,551 : Scope2 52,351을 앵커에 적용, 공장은 에너지 비율 (결정 19·27) | 가정: 직접 배출 월 = 생산 모양 (결정 17) |
| HBG_GJ | 같음 | 가정: 직접 배출 월 = 생산 모양 (결정 17) — 광주·송악은 회사 주력 제품 모양 (결정 40) |
| HBG_SA | 같음 | 같음 |
| SA_PS | 연료 먼저 — 직접 = 연료 TJ × 38.5 ÷ 42.7 × 56.1, 간접 = 앵커 − 직접 (결정 28·29) | 가정: 연료 = 도시가스, 명세서 에너지 = 총발열량, 직접 배출 월 = 생산 모양 (결정 17) |
| HI_CEM | 회사 사업별 공개값 — 시멘트 Scope1 5,712,566 / Scope2 370,165 (결정 20·27) | 가정: 직접 배출 월 = 생산 모양 (결정 17), 전력 월 = 가중 모양 (결정 21) |
| HI_SEC | 회사 사업별 공개값 — 2차제품 Scope1 54,947 / Scope2 40,418 (결정 20·27) | 가정: 직접 배출 월 = 생산 모양 (결정 17) |
"""
EMISSIONS_COLUMNS = ["month", "company_id", "plant_id", "direct_fuel_t", "process_t",
                     "direct_t", "indirect_t", "total_t", "direct_fuel_alt_net_t",
                     "split_basis", "direct_shape_id", "indirect_shape_id"] + COMMON_TAIL
EMISSIONS_NATURE = "합성"
EMISSIONS_METHOD = "연간 공장 값(4-8-7 ①) → 직접·공정은 제품 모양, 간접은 M1 모양으로 최대 나머지 배분 (0.001 t)"
EMISSIONS_SOURCE = "09 4-7-3·4-8-7"

PURCHASES_TABLE = """
| line_no | plant_id | material | eu_cn | default_cn_key | precursor | quantity_method | check_value | origin | supplier_note | transfer | zero_reason | shape_id | assumption |
|---|---|---|---|---|---|---|---|---|---|---|---|---|---|
| 1 | KG_DJ | 열연 코일 (HR COIL) | 7208 | 7208 | 해당 | 공시 금액 ÷ 공시 단가 1,509,441백만원 ÷ 739천원/t = 2,042,545 t (정수 반올림, 09-3) → 생산실적 비율 1,606 : 601 (결정 33) | 빈칸 | 불명 | 매입처·국내/수입 비율 미공시 | 빈칸 | 빈칸 | P_COLD | 가정: 원재료도 생산실적 비율로 공장에 들어감 (결정 33), 매입 월 = 생산 모양 (결정 41) |
| 2 | KG_DJ | 기타 코일 (BP·GI·SGL·AL 등) | 7209·7210 | 7210 | 해당 | 생산 − 열연 ÷ 1.04 (JRC134682 p.16) = 2,207,000 − 2,042,545 ÷ 1.04 = 243,014 t (정수 반올림, 09-3) → 1,606 : 601 (결정 33) | 내재 단가 220,135백만원 ÷ 243,014 t = 905,853원/t > 열연 739천원/t | 불명 | 매입처 미공시 | 빈칸 | 빈칸 | P_COLD | 가정: 기타 코일 → 제품 1:1 (JRC "a factor of 1 is assumed"), 공장 비율 (결정 33), 매입 월 = 생산 모양 (결정 41) |
| 3 | KG_IC | 열연 코일 (HR COIL) | 7208 | 7208 | 해당 | 공시 금액 ÷ 공시 단가 1,509,441백만원 ÷ 739천원/t = 2,042,545 t (정수 반올림, 09-3) → 생산실적 비율 1,606 : 601 (결정 33) | 빈칸 | 불명 | 매입처·국내/수입 비율 미공시 | 빈칸 | 빈칸 | P_COLD | 가정: 원재료도 생산실적 비율로 공장에 들어감 (결정 33), 매입 월 = 생산 모양 (결정 41) |
| 4 | KG_IC | 기타 코일 (BP·GI·SGL·AL 등) | 7209·7210 | 7210 | 해당 | 생산 − 열연 ÷ 1.04 (JRC134682 p.16) = 2,207,000 − 2,042,545 ÷ 1.04 = 243,014 t (정수 반올림, 09-3) → 1,606 : 601 (결정 33) | 내재 단가 220,135백만원 ÷ 243,014 t = 905,853원/t > 열연 739천원/t | 불명 | 매입처 미공시 | 빈칸 | 빈칸 | P_COLD | 가정: 기타 코일 → 제품 1:1 (JRC "a factor of 1 is assumed"), 공장 비율 (결정 33), 매입 월 = 생산 모양 (결정 41) |
| 5 | HBG_CW | 스테인리스 열연 강대 | 7219·7220 | 7219 | 해당 | 사 온 양 = 만든 양 (창원 생산 239,432 t, 부산물 포함 입고 기준) | 원재료와 소모품 사용액 485,894,535,877원 ÷ 2,158원/kg = 225,160 t (생산 대비 −6.0%) | 불명 | 포스코(국내)·주요 해외 원재사·현대머티리얼(계열), 비율 미공시 | 빈칸 | 빈칸 | P_COLD | 가정: 사 온 양 = 만든 양 (수율 1.0), 매입 월 = 생산 모양 (결정 41) |
| 6 | HBG_SA | 창원 판 (이동분) | 7219·7220 (추정) | 빈칸 | 해당 (추정) | 이동 수량 미공시 (송악 능력 3.1만 t/년만) | 빈칸 | 회사 안 (창원) | 현대비앤지스틸 창원 | 이동분 — 회사 합계에 더하지 않음 | 빈칸 | P_COLD | 빈칸 |
| 7 | TW_HJ | 철 스크랩 | 7204 | 빈칸 | 비해당 | 공시 금액 ÷ 공시 단가 141,443백만원 ÷ 677원/kg = 208,926 t (정수 반올림 — 09-3 4-6-2와 같음) | 빈칸 | 국내·해외 (비율 미공시) | 국내·해외 업체 | 빈칸 | 배출 0 — 철 스크랩 7204는 전구물질 아님 (2023/956 부속서 I 제외 품목) | P_STEEL1 | 가정: 매입 월 = 생산 모양 (결정 41) |
| 8 | TW_HJ | 철스크랩 외의 "외" | 불명 | 빈칸 | 불명 | 0 — 품목·수량 미공시 (결정 11) | 빈칸 | 불명 | 빈칸 | 빈칸 | 0 — 품목·수량 미공시, 합금철이면 일부 전구물질 (결정 11) | P_STEEL1 | 빈칸 |
| 9 | TW_NS | 화전 잉곳·블룸 (이동분) | 7206·7207·7218 등 (품목 추정) | TRANSFER:TW_HJ | 해당 | 받은 양 = 만든 양 — 녹산 단조 생산 68,165 t (결정 10 규칙) | 빈칸 | 회사 안 (화전) | 태웅 화전 (제강사업부) | 이동분 — 회사 합계에 더하지 않음 | 빈칸 | P_STEEL1 | 가정: 받은 양 = 만든 양, 매입 월 = 생산 모양 (결정 41) |
| 10 | TW_NS | 단조 원재료 외부 매입분 | 7206·7218 등 | 빈칸 | 해당 | 0 — 외부 매입 품목·수량 미공시 (원문 "제강사업부에서 주로 조달", 결정 16) | 빈칸 | 불명 | 빈칸 | 빈칸 | 0 — 미공시 (결정 16) | P_STEEL1 | 빈칸 |
| 11 | SA_PS | AL-STRIP | 7606 | 7606 | 해당 | 사 온 양 = 만든 양 (압연 15,457 t, 결정 12) | 상한 EA 2018 1.1587 → 17,910 t | 불명 | 주로 노벨리스코리아 (국내 매입), 생산국 증빙 없음 | 빈칸 | 빈칸 | P_ALFOIL | 가정: 사 온 양 = 만든 양 (1.0), 매입 월 = 생산 모양 (결정 41) |
| 12 | HI_CEM | 클링커 매입 | 2523 10 | 빈칸 | 해당 (매입 시) | 0 — 매입 없음, 자체 생산 7,430,004 t (결정 11) | 빈칸 | 빈칸 | 비상시 인근 시멘트사 임시 매입만 언급 (지속가능경영보고서) | 빈칸 | 0 — 자체 생산 (결정 11) | P_CEMENT | 빈칸 |
"""
PURCHASES_COLUMNS = ["month", "company_id", "plant_id", "line_no", "material", "eu_cn",
                     "default_cn_key", "precursor", "quantity_t", "quantity_method",
                     "check_value", "origin", "supplier_note", "transfer", "zero_reason",
                     "shape_id"] + COMMON_TAIL
PURCHASES_NATURE = "앵커 밖 합성"
PURCHASES_METHOD = "연간 수량 → 받는 공장 제품 모양으로 최대 나머지 배분 (0.001 t)"
PURCHASES_SOURCE = "09 4-6-2·4-6-3·4-7-4"

# 대조 칸(fuel_tj_net·direct_fuel_alt_net_t)과 연소·공정 구분을 두는 공장 — "연료 먼저"(결정 28·29)
FUEL_FIRST = {"KG_DJ", "KG_IC", "TW_HJ", "TW_NS", "SA_PS"}


def _md_rows(text):
    """마크다운 표(머리글 줄 + 구분 줄 + 값 줄)를 칸 이름 → 값 사전 목록으로 바꿈 (09 4-8-0 표 읽는 법)"""
    lines = [ln.strip() for ln in text.strip().splitlines()]
    head = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = []
    for ln in lines[2:]:
        cells = [c.strip() for c in ln.strip("|").split("|")]
        if len(cells) != len(head):
            raise ValueError(f"표 칸 수가 머리글과 다름: {ln}")
        if rows:
            cells = [rows[-1][i] if c == "같음" else c for i, c in enumerate(cells)]
        rows.append(cells)
    return [{h: (None if c == "빈칸" else c) for h, c in zip(head, r)} for r in rows]


MASTER = {r["plant_id"]: r for r in _md_rows(PLANT_MASTER_TABLE)}


def load_ext():
    """③ 파일 4개를 읽음"""
    anchor = {r["company_id"]: {"emissions": Fraction(r["emissions_tco2eq"]),
                                "energy": Fraction(r["energy_tj"])}
              for r in read_csv(EXT / "mfg_ext_anchor.csv")}

    wb = load_workbook(EXT / "mfg_ext_factors.xlsx", data_only=True)
    ws = wb["gov_factors"]
    it = ws.iter_rows(values_only=True)
    head = [str(c) for c in next(it)]
    gov = {}
    for row in it:
        r = dict(zip(head, row))
        gov[str(r["factor_id"])] = Fraction(str(r["value"]))

    stats = {}
    for r in read_csv(EXT / "mfg_ext_stats.csv"):
        if r["month"]:
            stats.setdefault(r["stat_id"], {})[r["month"]] = Fraction(r["value"])
        else:
            stats[r["stat_id"]] = Fraction(r["value"])
    for k, v in stats.items():
        if isinstance(v, dict):
            stats[k] = [v[m] for m in MONTHS]

    disc = {(r["company_id"], r["item_id"]): Fraction(r["value"])
            for r in read_csv(EXT / "mfg_ext_disclosed.csv")}
    return {"anchor": anchor, "gov": gov, "stats": stats, "disc": disc}


def build_shapes(ext):
    """월 모양(M1·제품·시멘트 가중)을 만듦 (09 4-8-4)"""
    st, dc = ext["stats"], ext["disc"]
    shapes = {}
    for sid, v in st.items():
        if isinstance(v, list) and (sid.startswith("M1_") or sid.startswith("P_")):
            total = sum(v)
            shapes[sid] = [x / total for x in v]
    # M1_CEMENT_MIX (결정 21): 충남 가중 = 슬래그 설비 ÷ 생산능력, 충북·강원은 나머지를 연간 kWh 비율로
    w_cn = dc[("HI", "SLAG_CAP")] / dc[("HI", "CAP_CEM")]
    cb, gw = sum(st["M1_CHUNGBUK"]), sum(st["M1_GANGWON"])
    w_cb = (1 - w_cn) * cb / (cb + gw)
    w_gw = 1 - w_cn - w_cb
    shapes["M1_CEMENT_MIX"] = [w_cb * a + w_gw * b + w_cn * c for a, b, c in
                               zip(shapes["M1_CHUNGBUK"], shapes["M1_GANGWON"],
                                   shapes["M1_CHUNGNAM"])]
    return shapes


def plant_annual(ext):
    """공장 연간 값을 계산함 (09 4-8-7 ①)"""
    an, gov, st, dc = ext["anchor"], ext["gov"], ext["stats"], ext["disc"]
    net = gov["LNG_NCV"] / gov["LNG_GCV"]
    gas = gov["GAS_EF"]
    mj = gov["ELEC_MJ_PER_KWH"]

    def e(code):
        # 전기 에너지 몫 = (전력비 ÷ 전력 단가 × 9.6) ÷ (전력비 ÷ 전력 단가 × 9.6 + 연료비 ÷ 가스 단가)
        elec = st[f"M2_{code}_ELEC"] / st["PRICE_ELEC_2024"] * mj
        fuel = st[f"M2_{code}_FUEL"] / st["PRICE_GAS_2024"]
        return elec / (elec + fuel)

    out = {}

    # KG: 회사 연료·전기·직접·간접 → 생산실적 비율로 당진·인천 (결정 1·28·29)
    en, em = an["KG"]["energy"], an["KG"]["emissions"]
    fuel = r3(en * (1 - e("24122")))
    elec = en - fuel
    direct = r3(fuel * net * gas)
    indirect = em - direct
    w = [dc[("KG", "PROD_DJ")], dc[("KG", "PROD_IC")]]
    for pid, f, el, d, i in zip(["KG_DJ", "KG_IC"], lr(fuel, w), lr(elec, w), lr(direct, w),
                                lr(indirect, w)):
        out[pid] = {"fuel": f, "elec": el, "direct_fuel": d, "process": Fraction(0),
                    "direct": None, "indirect": i}

    # TW: 업종 원단위 × 사업부 금액으로 에너지 → 화전·녹산 (결정 18·28)
    en, em = an["TW"]["energy"], an["TW"]["emissions"]
    ens = lr(en, [dc[("TW", "FORGE_RAW_USE")] * st["INT_24112"],
                  dc[("TW", "FORGE_SALES")] * st["INT_25912"]])
    fuels = [r3(ens[0] * (1 - e("24112"))), r3(ens[1] * (1 - e("25912")))]
    elecs = [x - f for x, f in zip(ens, fuels)]
    dfs = [r3(f * net * gas) for f in fuels]
    procs = [dc[("TW", "PROD_STEEL")] * gov["EAF_EF"], Fraction(0)]
    indirect = em - sum(dfs) - sum(procs)
    for pid, f, el, d, p, i in zip(["TW_HJ", "TW_NS"], fuels, elecs, dfs, procs,
                                   lr(indirect, elecs)):
        out[pid] = {"fuel": f, "elec": el, "direct_fuel": d, "process": p,
                    "direct": None, "indirect": i}

    # HBG: 광주·송악은 업종 원단위, 창원은 나머지 (결정 19·27)
    en, em = an["HBG"]["energy"], an["HBG"]["emissions"]
    e_gj = r3(dc[("HBG", "SALES_GJ")] * st["INT_30310"])
    e_sa = r3(dc[("HBG", "PROC_REV_SA")] * st["INT_24122"])
    ens = [en - e_gj - e_sa, e_gj, e_sa]
    f_tj, e_tj = dc[("HBG", "FUEL_TJ")], dc[("HBG", "ELEC_TJ")]
    fuels = [r3(x * f_tj / (f_tj + e_tj)) for x in ens]
    s1, s2 = dc[("HBG", "S1")], dc[("HBG", "S2")]
    direct = r3(em * s1 / (s1 + s2))
    indirect = em - direct
    for pid, x, f, d, i in zip(["HBG_CW", "HBG_GJ", "HBG_SA"], ens, fuels, lr(direct, ens),
                               lr(indirect, ens)):
        out[pid] = {"fuel": f, "elec": x - f, "direct_fuel": None, "process": None,
                    "direct": d, "indirect": i}

    # SA: 연료 먼저 (결정 28·29)
    en, em = an["SA"]["energy"], an["SA"]["emissions"]
    fuel = r3(en * (1 - e("24222")))
    direct = r3(fuel * net * gas)
    out["SA_PS"] = {"fuel": fuel, "elec": en - fuel, "direct_fuel": direct,
                    "process": Fraction(0), "direct": None, "indirect": em - direct}

    # HI: 묶음 배출 = 사업별 공개값, 묶음 에너지 = 결정 39
    elecs = lr(dc[("HI", "ELEC_TJ")], [dc[("HI", "CEM_S2")], dc[("HI", "SEC_S2")]])
    fuel_sec = r3(dc[("HI", "SEC_S1")] / gas / net)
    fuel_cem = (an["HI"]["energy"] - dc[("HI", "ELEC_TJ")]) - fuel_sec
    out["HI_CEM"] = {"fuel": fuel_cem, "elec": elecs[0], "direct_fuel": None, "process": None,
                     "direct": dc[("HI", "CEM_S1")], "indirect": dc[("HI", "CEM_S2")]}
    out["HI_SEC"] = {"fuel": fuel_sec, "elec": elecs[1], "direct_fuel": None, "process": None,
                     "direct": dc[("HI", "SEC_S1")], "indirect": dc[("HI", "SEC_S2")]}
    return out


def monthly_values(annual, shapes, ext):
    """공장 연간 값을 달로 나눔 (09 4-8-7 ②)"""
    gov = ext["gov"]
    net = gov["LNG_NCV"] / gov["LNG_GCV"]
    gas = gov["GAS_EF"]
    mj = gov["ELEC_MJ_PER_KWH"]
    blank = [None] * 12
    out = {}
    for pid in PLANTS:
        a = annual[pid]
        es = shapes[MASTER[pid]["elec_shape_id"]]
        ps = shapes[MASTER[pid]["prod_shape_id"]]
        m = {"fuel": lr(a["fuel"], ps), "elec": lr(a["elec"], es),
             "indirect": lr(a["indirect"], es)}
        if pid in FUEL_FIRST:
            m["direct_fuel"] = lr(a["direct_fuel"], ps)
            m["process"] = lr(a["process"], ps)
            m["direct"] = [d + p for d, p in zip(m["direct_fuel"], m["process"])]
            m["fuel_net"] = [r3(x * net) for x in m["fuel"]]
            m["alt_net"] = [r3(x * gas) for x in m["fuel"]]
        else:
            m["direct_fuel"] = m["process"] = m["fuel_net"] = m["alt_net"] = blank
            m["direct"] = lr(a["direct"], ps)
        m["total"] = [d + i for d, i in zip(m["direct"], m["indirect"])]
        m["energy"] = [f + el for f, el in zip(m["fuel"], m["elec"])]
        m["mwh"] = [r3(x * 1000 / mj) for x in m["elec"]]
        out[pid] = m
    return out


def build_plant_master(created):
    """공장 목록 행을 만듦 (09 4-8-6)"""
    return [{**r, "data_nature": PLANT_MASTER_NATURE, "method": PLANT_MASTER_METHOD,
             "source": PLANT_MASTER_SOURCE, "created_at": created, "version": VERSION}
            for r in _md_rows(PLANT_MASTER_TABLE)]


def build_energy(monthly, created):
    """월별 에너지 행을 만듦 (09 4-8-6)"""
    text = {r["plant_id"]: r for r in _md_rows(ENERGY_TABLE)}
    rows = []
    for pid in PLANTS:
        m, t = monthly[pid], text[pid]
        for k, month in enumerate(MONTHS):
            rows.append({"month": month, "company_id": MASTER[pid]["company_id"],
                         "plant_id": pid, "fuel_tj_gross": m["fuel"][k],
                         "fuel_tj_net": m["fuel_net"][k], "elec_tj": m["elec"][k],
                         "elec_mwh": m["mwh"][k], "energy_tj": m["energy"][k],
                         "ratio_basis": t["ratio_basis"],
                         "fuel_shape_id": MASTER[pid]["prod_shape_id"],
                         "elec_shape_id": MASTER[pid]["elec_shape_id"],
                         "envizi_input": ENERGY_ENVIZI, "data_nature": ENERGY_NATURE,
                         "method": ENERGY_METHOD, "assumption": t["assumption"],
                         "source": ENERGY_SOURCE, "created_at": created, "version": VERSION})
    return rows


def build_production(ext, shapes, created):
    """월별 생산 행을 만듦 (09 4-8-6)"""
    dc = ext["disc"]
    rows = []
    for r in _md_rows(PRODUCTION_TABLE):
        item = PRODUCTION_ITEM[r["line_no"]]
        if item is None:
            annual, quantities, method = None, [None] * 12, PRODUCTION_METHOD_BLANK
        else:
            cid, iid, scale = item
            annual = dc[(cid, iid)] * scale
            unit = 1 if r["unit"] == "개" else Fraction(1, 1000)
            quantities = lr(annual, shapes[r["prod_shape_id"]], unit)
            method = PRODUCTION_METHOD
            if r["unit"] == "개":
                annual, quantities = int(annual), [int(q) for q in quantities]
        for k, month in enumerate(MONTHS):
            rows.append({"month": month, "company_id": MASTER[r["plant_id"]]["company_id"],
                         "plant_id": r["plant_id"], "line_no": int(r["line_no"]),
                         "product": r["product"], "eu_cn": r["eu_cn"],
                         "quantity": quantities[k], "unit": r["unit"],
                         "annual_public": annual, "prod_shape_id": r["prod_shape_id"],
                         "data_nature": PRODUCTION_NATURE, "method": method,
                         "assumption": r["assumption"], "source": r["source"],
                         "created_at": created, "version": VERSION})
    return rows


def build_emissions(monthly, created):
    """월별 배출 행을 만듦 (09 4-8-6)"""
    text = {r["plant_id"]: r for r in _md_rows(EMISSIONS_TABLE)}
    rows = []
    for pid in PLANTS:
        m, t = monthly[pid], text[pid]
        for k, month in enumerate(MONTHS):
            rows.append({"month": month, "company_id": MASTER[pid]["company_id"],
                         "plant_id": pid, "direct_fuel_t": m["direct_fuel"][k],
                         "process_t": m["process"][k], "direct_t": m["direct"][k],
                         "indirect_t": m["indirect"][k], "total_t": m["total"][k],
                         "direct_fuel_alt_net_t": m["alt_net"][k],
                         "split_basis": t["split_basis"],
                         "direct_shape_id": MASTER[pid]["prod_shape_id"],
                         "indirect_shape_id": MASTER[pid]["elec_shape_id"],
                         "data_nature": EMISSIONS_NATURE, "method": EMISSIONS_METHOD,
                         "assumption": t["assumption"], "source": EMISSIONS_SOURCE,
                         "created_at": created, "version": VERSION})
    return rows


def build_purchases(ext, shapes, created):
    """월별 원재료 매입 행을 만듦 (09 4-8-6)"""
    st, dc = ext["stats"], ext["disc"]
    # 연간 수량 (09 4-8-6 연간 값 계산) — 금액·단가 단위 바꾸기: 백만원 × 10⁶, 천원 × 10³, 원/kg ÷ 1,000 → t
    hr = r0(dc[("KG", "HR_AMT")] * 10 ** 6 / (dc[("KG", "HR_PRICE")] * 10 ** 3))
    kg_prod = (dc[("KG", "PROD_DJ")] + dc[("KG", "PROD_IC")]) * 1000
    other = r0(kg_prod - hr / st["JRC_HR_PER_CR"])
    w = [dc[("KG", "PROD_DJ")], dc[("KG", "PROD_IC")]]
    hr_split, other_split = lr(hr, w), lr(other, w)
    scrap = r0(dc[("TW", "SCRAP_AMT")] * 10 ** 6 / dc[("TW", "SCRAP_PRICE")] / 1000)
    annual = {"1": hr_split[0], "2": other_split[0], "3": hr_split[1], "4": other_split[1],
              "5": dc[("HBG", "PROD_CW")], "6": None, "7": scrap, "8": Fraction(0),
              "9": dc[("TW", "PROD_FORGE")], "10": Fraction(0), "11": dc[("SA", "PROD_ROLL")],
              "12": Fraction(0)}
    rows = []
    for r in _md_rows(PURCHASES_TABLE):
        a = annual[r["line_no"]]
        quantities = [None] * 12 if a is None else lr(a, shapes[r["shape_id"]])
        for k, month in enumerate(MONTHS):
            rows.append({"month": month, "company_id": MASTER[r["plant_id"]]["company_id"],
                         "plant_id": r["plant_id"], "line_no": int(r["line_no"]),
                         "material": r["material"], "eu_cn": r["eu_cn"],
                         "default_cn_key": r["default_cn_key"], "precursor": r["precursor"],
                         "quantity_t": quantities[k], "quantity_method": r["quantity_method"],
                         "check_value": r["check_value"], "origin": r["origin"],
                         "supplier_note": r["supplier_note"], "transfer": r["transfer"],
                         "zero_reason": r["zero_reason"], "shape_id": r["shape_id"],
                         "data_nature": PURCHASES_NATURE, "method": PURCHASES_METHOD,
                         "assumption": r["assumption"], "source": PURCHASES_SOURCE,
                         "created_at": created, "version": VERSION})
    return rows


def main():
    """① 파일 5개를 씀"""
    created = today_str()
    ext = load_ext()
    shapes = build_shapes(ext)
    monthly = monthly_values(plant_annual(ext), shapes, ext)
    outputs = [
        ("mfg_plant_master_synth.csv", PLANT_MASTER_COLUMNS, build_plant_master(created)),
        ("mfg_energy_monthly_synth.csv", ENERGY_COLUMNS, build_energy(monthly, created)),
        ("mfg_production_monthly_synth.csv", PRODUCTION_COLUMNS,
         build_production(ext, shapes, created)),
        ("mfg_emissions_monthly_synth.csv", EMISSIONS_COLUMNS, build_emissions(monthly, created)),
        ("mfg_purchases_monthly_synth.csv", PURCHASES_COLUMNS,
         build_purchases(ext, shapes, created)),
    ]
    for name, columns, rows in outputs:
        write_csv(DB / name, columns, rows)
        print(f"{name}: {len(rows)}행")


if __name__ == "__main__":
    main()
